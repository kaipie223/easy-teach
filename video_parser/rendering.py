from __future__ import annotations

import html
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from .intermediate_schemas import (
    ArtifactFile,
    ArtifactManifest,
    DemoGenerationPlan,
    InteractiveSpec,
    LessonPlanSpec,
    SlideDeckSpec,
)
from .package import LoadedTeachingContentPackage, file_sha256
from .planning import build_demo_generation_plan, build_interactive_spec, build_lesson_plan_spec, build_slide_deck_spec
from .quality import evaluate_generation


class RenderingError(RuntimeError):
    pass


# 渲染产物的完整清单；失败时按它清理，避免把半个产物目录留给消费方。
_RENDER_ARTIFACT_NAMES = (
    "generation_plan.json",
    "slide_spec.json",
    "lesson_plan_spec.json",
    "interactive_spec.json",
    "quality_report.json",
    "generation_access_audit.json",
    "teaching_demo.pptx",
    "lesson_plan.docx",
    "content_summary.md",
    "preview.svg",
    "artifact_manifest.json",
)


def _cleanup_failed_render(root: Path) -> None:
    """清掉本次写出的产物，以及上一次成功运行留下的 artifact_manifest.json。

    manifest 记录的是上一次成功产物的哈希；失败后若还留着它，消费方按它校验就会
    拿到 MISMATCH，且无法区分"下载损坏"和"刚刚生成失败"。
    """
    for name in _RENDER_ARTIFACT_NAMES:
        (root / name).unlink(missing_ok=True)
    shutil.rmtree(root / "interactive_html", ignore_errors=True)


def render_demo_outputs(
    package: LoadedTeachingContentPackage,
    output_dir: str | Path,
    *,
    plan: DemoGenerationPlan | None = None,
    run_label: str = "candidate",
) -> dict[str, Any]:
    """Render all demo outputs from a package without consulting the source video."""
    root = Path(output_dir).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    plan = plan or build_demo_generation_plan(package)
    slides = build_slide_deck_spec(package, plan)
    lesson = build_lesson_plan_spec(package, plan)
    interactive = build_interactive_spec(package, plan)
    # 质量门禁先于任何落盘：门禁失败时既不该留下半个产物，也不该留下上一次成功
    # 运行写下的 artifact_manifest.json —— 否则消费方按它校验会拿到 MISMATCH，
    # 且无法区分"下载损坏"和"刚刚生成失败"。
    quality = evaluate_generation(package, plan, slides, lesson, interactive, run_label=run_label)
    if quality.errors:
        _cleanup_failed_render(root)
        raise RenderingError("Quality gate blocked export: " + "; ".join(quality.errors[:3]))

    _write_json(root / "generation_plan.json", plan.model_dump(mode="json"))
    _write_json(root / "slide_spec.json", slides.model_dump(mode="json"))
    _write_json(root / "lesson_plan_spec.json", lesson.model_dump(mode="json"))
    _write_json(root / "interactive_spec.json", interactive.model_dump(mode="json"))

    try:
        pptx_path = root / "teaching_demo.pptx"
        docx_path = root / "lesson_plan.docx"
        html_dir = root / "interactive_html"
        _render_pptx(package, slides, pptx_path)
        _render_docx(package, lesson, docx_path)
        _render_html(package, interactive, html_dir)
        summary_path = root / "content_summary.md"
        _render_markdown_summary(package, plan, summary_path)
        preview_path = root / "preview.svg"
        _render_preview(slides, preview_path)
        _write_json(root / "quality_report.json", quality.model_dump(mode="json"))
        audit = {
            "stage": "generation",
            "read_files": package.audit_snapshot(),
            "source_video_reads": 0,
            "source_video_access_required": False,
            "only_package_relative_reads": all(_is_package_relative(path) for path in package.audit_snapshot()),
        }
        _write_json(root / "generation_access_audit.json", audit)

        artifact_files = []
        for path, artifact_type in [
            (root / "generation_plan.json", "other"),
            (root / "slide_spec.json", "other"),
            (root / "lesson_plan_spec.json", "other"),
            (root / "interactive_spec.json", "other"),
            (pptx_path, "pptx"),
            (docx_path, "docx"),
            (preview_path, "preview"),
            (root / "quality_report.json", "quality"),
            (root / "generation_access_audit.json", "quality"),
            (summary_path, "other"),
        ]:
            artifact_files.append(
                ArtifactFile(
                    path=path.relative_to(root).as_posix(),
                    sha256=file_sha256(path),
                    size_bytes=path.stat().st_size,
                    artifact_type=artifact_type,  # type: ignore[arg-type]
                )
            )
        # pptx_previews 只可能由已删除的 artifact-tool 路线产生，该目录永不出现，
        # 这段收集逻辑是恒假分支，一并移除。
        for path in html_dir.rglob("*"):
            if path.is_file():
                artifact_files.append(
                    ArtifactFile(
                        path=path.relative_to(root).as_posix(),
                        sha256=file_sha256(path),
                        size_bytes=path.stat().st_size,
                        artifact_type="html",
                    )
                )
        manifest = ArtifactManifest(
            plan_id=plan.plan_id,
            package_id=package.manifest.package_id,
            package_version=package.manifest.package_version,
            files=artifact_files,
            source_refs=plan.source_refs,
            warnings=quality.warnings,
            status="complete" if quality.status != "error" else "partial",
        )
        _write_json(root / "artifact_manifest.json", manifest.model_dump(mode="json"))
    except Exception:
        # 渲染中途失败同样不能留下半个产物目录（含上一次成功留下的陈旧 manifest）。
        _cleanup_failed_render(root)
        raise
    return {
        "output_dir": root,
        "manifest": manifest,
        "quality": quality,
        "plan": plan,
        "slides": slides,
        "lesson": lesson,
        "interactive": interactive,
        "audit": audit,
    }


def _render_pptx(package: LoadedTeachingContentPackage, spec: SlideDeckSpec, path: Path) -> None:
    """Render the deck with the maintained legacy renderer.

    An "artifact-tool" JS route used to be attempted first, but the script it
    invoked has never existed in this repository and its default node/setup
    paths pointed at another developer's machine, so it always failed and was
    silently swallowed.  The dead route is removed; PPTX rendering is now
    predictable and always uses the working renderer.
    """
    _render_pptx_legacy(package, spec, path)


def _render_pptx_legacy(package: LoadedTeachingContentPackage, spec: SlideDeckSpec, path: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RenderingError("PPTX renderer requires python-pptx") from exc

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    background = RGBColor(248, 247, 243)
    ink = RGBColor(24, 35, 52)
    accent = RGBColor(28, 94, 112)
    muted = RGBColor(102, 112, 123)
    for slide_spec in spec.slides:
        slide = presentation.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = background
        band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
        band.fill.solid()
        band.fill.fore_color.rgb = accent
        band.line.fill.background()
        title = slide.shapes.add_textbox(Inches(0.65), Inches(0.52), Inches(12), Inches(0.65))
        title_frame = title.text_frame
        title_frame.clear()
        title_frame.margin_left = 0
        paragraph = title_frame.paragraphs[0]
        paragraph.text = _truncate(slide_spec.title, 72)
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(50 if slide_spec.role == "title" else 35)
        paragraph.font.bold = True
        paragraph.font.color.rgb = ink
        content_top = 1.5 if slide_spec.role != "title" else 2.05
        y = content_top
        image_index = 0
        for element in slide_spec.elements:
            if element.visibility == "hidden":
                continue
            if element.element_type == "image" and element.asset_refs:
                try:
                    asset_path = package.asset_path(element.asset_refs[0])
                    image_top = 1.7 + image_index * 2.55
                    if image_top + 2.35 <= 6.85 and image_index < 2:
                        slide.shapes.add_picture(
                            str(asset_path),
                            Inches(7.6),
                            Inches(image_top),
                            width=Inches(4.9),
                            height=Inches(2.35),
                        )
                        image_index += 1
                    continue
                except (OSError, ValueError):
                    pass
            text = element.text or (f"公式：{element.latex}" if element.latex else "")
            if element.element_type == "bullet":
                text = f"• {text}"
            if element.element_type == "source":
                text = f"来源：{text}"
            if element.element_type == "answer":
                text = f"参考答案：{text}"
            font_size = 22 if element.element_type == "formula" else 18
            box_height = _estimate_pptx_text_box_height(text, font_size)
            box = slide.shapes.add_textbox(Inches(0.85), Inches(y), Inches(6.25), Inches(box_height))
            frame = box.text_frame
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.TOP
            frame.margin_left = 0
            frame.margin_right = 0
            paragraph = frame.paragraphs[0]
            paragraph.text = _truncate(text, 360)
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(font_size)
            paragraph.font.color.rgb = ink if element.visibility != "teacher" else muted
            if element.element_type in {"formula", "question", "answer"}:
                paragraph.font.bold = True
            y += box_height + 0.16
            if y > 6.45:
                break
        footer = slide.shapes.add_textbox(Inches(0.85), Inches(7.05), Inches(11.8), Inches(0.25))
        footer_frame = footer.text_frame
        footer_frame.clear()
        footer_frame.paragraphs[0].text = f"{slide_spec.order:02d}  ·  证据 {len(slide_spec.source_refs)} 条"
        footer_frame.paragraphs[0].font.name = "Microsoft YaHei"
        footer_frame.paragraphs[0].font.size = Pt(9)
        footer_frame.paragraphs[0].font.color.rgb = muted
        footer_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    presentation.save(path)


_DOCX_LABELS = {
    "lesson_plan": "\u6559\u5b66\u65b9\u6848",
    "audience": "\u5bf9\u8c61",
    "preset": "\u6587\u6863\u9884\u8bbe",
    "objectives": "\u6559\u5b66\u76ee\u6807",
    "key_points": "\u91cd\u70b9\u5185\u5bb9",
    "difficulties": "\u91cd\u70b9\u96be\u70b9",
    "preparation": "\u8bfe\u524d\u51c6\u5907",
    "process": "\u6559\u5b66\u8fc7\u7a0b",
    "time": "\u65f6\u95f4",
    "minutes": "\u5206\u949f",
    "teacher": "\u6559\u5e08\u6d3b\u52a8",
    "student": "\u5b66\u751f\u6d3b\u52a8",
    "assessment": "\u8bc4\u4ef7",
    "evidence": "\u8bc1\u636e\u56de\u6eaf",
    "homework": "\u8bfe\u540e\u4f5c\u4e1a",
    "reflection": "\u6559\u5b66\u53cd\u601d",
    "teacher_notes": "\u6559\u5e08\u590d\u6838\u63d0\u793a",
    "sources": "\u6765\u6e90\u56de\u6eaf",
    "none": "\u6682\u65e0\u5df2\u6807\u8bb0\u5185\u5bb9",
    "trace_note": "\u5b8c\u6574 Evidence ID \u4e0e\u539f\u59cb\u8bc1\u636e\u4fdd\u5b58\u5728\u4e2d\u95f4\u5305\u7684 generation_plan.json \u548c artifact_manifest.json\u4e2d\u3002",
}


def _render_docx_compact(package: LoadedTeachingContentPackage, spec: LessonPlanSpec, path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.shared import Pt
    except ImportError as exc:
        # 缺 python-docx 必须转成可识别的 RenderingError：裸 ImportError 会让调用方
        # 只看到"PPTX 已写出、DOCX 缺失"的半个产物，却判断不出原因。
        raise RenderingError("DOCX renderer requires python-docx") from exc

    document = Document()
    _configure_docx_styles(document)
    _configure_page(document)
    _configure_header_footer(document, spec.title)

    title = document.add_paragraph(style="Document Title")
    title.add_run(spec.title or _DOCX_LABELS["lesson_plan"])
    meta = document.add_paragraph(style="Document Meta")
    meta.add_run(
        f"{_DOCX_LABELS['audience']}：{spec.audience}  ·  "
        f"{_DOCX_LABELS['preset']}：compact_reference_guide"
    )
    _add_rule(meta)

    _add_compact_bullets(document, _DOCX_LABELS["objectives"], spec.objectives)
    _add_compact_bullets(document, _DOCX_LABELS["key_points"], spec.key_points)
    _add_compact_bullets(document, _DOCX_LABELS["difficulties"], spec.difficulties or [_DOCX_LABELS["none"]])
    _add_compact_bullets(document, _DOCX_LABELS["preparation"], spec.preparation)

    document.add_heading(_DOCX_LABELS["process"], level=1)
    for section in spec.sections:
        heading = document.add_heading(section.title, level=2)
        heading.paragraph_format.keep_with_next = True
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        header_cells = table.add_row().cells
        _set_cell_lines(header_cells[0], ["项目"], "Table Label")
        _set_cell_lines(header_cells[1], ["内容"], "Table Body")
        _set_cell_shading(header_cells[0], "E8EEF5")
        _set_cell_shading(header_cells[1], "E8EEF5")
        _mark_table_header(table.rows[0])
        rows = [
            (_DOCX_LABELS["time"], [f"{section.minutes} {_DOCX_LABELS['minutes']}"]),
            (_DOCX_LABELS["teacher"], section.teacher_activity),
            (_DOCX_LABELS["student"], section.student_activity),
            (_DOCX_LABELS["assessment"], section.assessment),
            (_DOCX_LABELS["evidence"], [_evidence_summary(section.source_refs)]),
        ]
        for label, values in rows:
            cells = table.add_row().cells
            _set_cell_lines(cells[0], [label], "Table Label")
            _set_cell_lines(cells[1], values or [_DOCX_LABELS["none"]], "Table Body")
            _set_cell_shading(cells[0], "F2F4F7")
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_table_geometry(table, [1701, 7659])
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(2)

    _add_compact_bullets(document, _DOCX_LABELS["homework"], spec.homework)
    _add_compact_bullets(document, _DOCX_LABELS["reflection"], spec.reflection)
    if spec.teacher_only_notes:
        _add_compact_bullets(document, _DOCX_LABELS["teacher_notes"], spec.teacher_only_notes)
    document.add_heading(_DOCX_LABELS["sources"], level=1)
    document.add_paragraph(_evidence_summary(spec.source_refs, limit=16), style="Document Meta")
    document.add_paragraph(_DOCX_LABELS["trace_note"], style="Document Meta")
    document.save(path)


def _configure_docx_styles(document) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    styles = document.styles

    def set_font(style, name="Calibri", size=11, bold=False, color="000000", east_asia="Microsoft YaHei"):
        style.font.name = name
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        r_fonts.set(qn("w:ascii"), name)
        r_fonts.set(qn("w:hAnsi"), name)
        r_fonts.set(qn("w:eastAsia"), east_asia)

    normal = styles["Normal"]
    set_font(normal, size=11, color="000000")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = 0

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        set_font(style, size=size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True

    list_style = styles["List Bullet"]
    set_font(list_style, size=11, color="000000")
    list_style.paragraph_format.left_indent = Inches(0.375)
    list_style.paragraph_format.first_line_indent = Inches(-0.188)
    list_style.paragraph_format.space_before = Pt(0)
    list_style.paragraph_format.space_after = Pt(4)
    list_style.paragraph_format.line_spacing = 1.25

    def add_style(name, size, color, bold=False, after=0, before=0, line_spacing=1.15):
        try:
            style = styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_font(style, size=size, bold=bold, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line_spacing
        return style

    title_style = add_style("Document Title", 23, "0B2545", bold=True, after=4, line_spacing=1.0)
    title_style.paragraph_format.keep_with_next = True
    add_style("Document Meta", 9.5, "5B6470", after=5, line_spacing=1.15)
    add_style("Table Label", 10, "1F3A5F", bold=True, after=2, line_spacing=1.15)
    add_style("Table Body", 10, "000000", after=2, line_spacing=1.15)


def _configure_page(document) -> None:
    from docx.shared import Inches

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def _configure_header_footer(document, title: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for section in document.sections:
        header = section.header
        header_paragraph = header.paragraphs[0]
        header_paragraph.text = ""
        header_paragraph.style = "Document Meta"
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_paragraph.add_run(title or _DOCX_LABELS["lesson_plan"])

        footer = section.footer
        footer_paragraph = footer.paragraphs[0]
        footer_paragraph.text = ""
        footer_paragraph.style = "Document Meta"
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_paragraph.add_run("TeachingContentPackage  ·  ")
        _add_page_field(footer_paragraph, OxmlElement, qn)


def _add_rule(paragraph, color="3D8DFF") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def _add_page_field(paragraph, OxmlElement, qn) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, visible, end])


def _add_compact_bullets(document, heading: str, values: list[str]) -> None:
    document.add_heading(heading, level=1)
    for value in values or [_DOCX_LABELS["none"]]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(str(value).strip())


def _set_cell_lines(cell, values: list[str], style_name: str) -> None:
    cell.text = ""
    for index, value in enumerate(values or [_DOCX_LABELS["none"]]):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.style = style_name
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 2
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.add_run(str(value).strip())


def _mark_table_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_table_geometry(table, widths: list[int]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    total = sum(widths)
    table.autofit = False
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table._tbl.tblPr.append(table_width)
    table_width.set(qn("w:w"), str(total))
    table_width.set(qn("w:type"), "dxa")

    indent = table._tbl.tblPr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table._tbl.tblPr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    layout = table._tbl.tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table._tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "CBD5E1")

    cell_margins = table._tbl.tblPr.find(qn("w:tblCellMar"))
    if cell_margins is None:
        cell_margins = OxmlElement("w:tblCellMar")
        table._tbl.tblPr.append(cell_margins)
    for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        margin = cell_margins.find(qn(f"w:{edge}"))
        if margin is None:
            margin = OxmlElement(f"w:{edge}")
            cell_margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            cell_width = tc_pr.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                tc_pr.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _evidence_summary(refs: list[str], limit: int = 8) -> str:
    unique = list(dict.fromkeys(str(ref) for ref in refs if ref))
    if not unique:
        return _DOCX_LABELS["none"]
    sample = ", ".join(unique[:limit])
    remainder = f"，另有 {len(unique) - limit} 条" if len(unique) > limit else ""
    return f"共 {len(unique)} 条：{sample}{remainder}。"


def _render_docx(package: LoadedTeachingContentPackage, spec: LessonPlanSpec, path: Path) -> None:
    # 早先这里在 return 之后还留了一整套 docx 渲染实现，永远执行不到，
    # 却让人误以为"缺 python-docx 会被转成 RenderingError"。死分支已删除，
    # 依赖缺失现在由 _render_docx_compact 顶部统一转换。
    _render_docx_compact(package, spec, path)


def _render_html(package: LoadedTeachingContentPackage, spec: InteractiveSpec, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(spec.model_dump(mode="json"), ensure_ascii=False).replace("</", "<\\/")
    title = html.escape(spec.title)
    document = f'''<!doctype html>
<html lang="zh-CN">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title>
<style>
:root {{ color-scheme: light; font-family: "Microsoft YaHei", system-ui, sans-serif; background:#f5f7f6; color:#182334; }}
body {{ margin:0; padding:32px 16px; }}
main {{ max-width:820px; margin:0 auto; }}
header {{ border-top:5px solid #1c5e70; padding:18px 0 12px; }}
h1 {{ margin:0 0 8px; font-size:clamp(1.65rem,4vw,2.35rem); }}
.objective {{ color:#52606d; margin:0 0 20px; }}
.progress {{ height:8px; background:#dce4e3; border-radius:4px; overflow:hidden; }}
.progress i {{ display:block; height:100%; width:0; background:#e49a3a; transition:width .2s ease; }}
.card {{ background:white; border:1px solid #d8e0df; border-radius:8px; padding:20px; margin:18px 0; box-shadow:0 4px 16px #1823340d; }}
.card h2 {{ margin-top:0; font-size:1.2rem; }}
button, input {{ font:inherit; }}
button {{ background:#1c5e70; color:white; border:0; border-radius:5px; padding:10px 16px; cursor:pointer; }}
button:focus-visible, input:focus-visible {{ outline:3px solid #e49a3a; outline-offset:2px; }}
label {{ display:block; padding:8px 0; }}
.feedback {{ min-height:1.5em; margin-top:12px; font-weight:600; }}
.feedback.good {{ color:#18704b; }} .feedback.retry {{ color:#a34b22; }}
.source {{ color:#66737f; font-size:.85rem; }}
</style></head>
<body><main><header><h1>{title}</h1><p id="objective" class="objective"></p><div class="progress" aria-label="完成进度"><i id="bar"></i></div><p id="progressText" class="source" aria-live="polite"></p></header><section id="app"></section></main>
<script>
const SPEC = {payload};
const app = document.getElementById('app');
const objective = document.getElementById('objective');
const bar = document.getElementById('bar');
const progressText = document.getElementById('progressText');
objective.textContent = SPEC.objective;
let completed = 0;
function esc(value) {{ return String(value ?? '').replace(/[&<>"']/g, char => ({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[char])); }}
function updateProgress() {{ const total = SPEC.questions.length; const percent = total ? Math.round(completed / total * 100) : 100; bar.style.width = percent + '%'; progressText.textContent = `已完成 ${{completed}} / ${{total}}`; }}
function render() {{
  if (!SPEC.questions.length) {{ app.innerHTML = '<div class="card"><h2>暂无互动题目</h2><p>请依据来源关键帧完成口头复述，并记录需要复核的内容。</p></div>'; updateProgress(); return; }}
  app.innerHTML = SPEC.questions.map((question, index) => {{
    const options = question.options?.length ? '<fieldset><legend class="source">请选择一个答案</legend>' + question.options.map(option => `<label><input type="radio" name="q-${{index}}" value="${{esc(option)}}"> ${{esc(option)}}</label>`).join('') + '</fieldset>' : `<label for="answer-${{index}}" class="source">写下你的回答</label><input id="answer-${{index}}" type="text" style="width:100%;box-sizing:border-box;padding:9px;border:1px solid #bcc8c6;border-radius:4px">`;
    return `<article class="card" data-index="${{index}}"><h2>${{index + 1}}. ${{esc(question.prompt)}}</h2>${{options}}<button type="button" data-check="${{index}}">检查回答</button><div class="feedback" id="feedback-${{index}}" aria-live="polite"></div><p class="source">证据：${{esc((question.source_refs || []).join(', ') || '未指定')}}</p></article>`;
  }}).join('');
  app.querySelectorAll('[data-check]').forEach(button => button.addEventListener('click', () => check(Number(button.dataset.check))));
  updateProgress();
}}
function check(index) {{
  const question = SPEC.questions[index]; const card = app.querySelector(`[data-index="${{index}}"]`); const feedback = card.querySelector('.feedback'); let answer = '';
  if (question.options?.length) {{ answer = card.querySelector('input:checked')?.value || ''; }} else {{ answer = card.querySelector('input[type=text]')?.value.trim() || ''; }}
  if (!answer) {{ feedback.textContent = '请先完成回答。'; feedback.className = 'feedback retry'; return; }}
  const correct = question.question_type === 'single_choice' && answer === question.correct_answer;
  if (question.question_type === 'reflection' || question.question_type === 'open_text') {{ feedback.textContent = question.answer_explanation || '请对照来源证据自检。'; feedback.className = 'feedback good'; }}
  else if (correct) {{ feedback.textContent = question.feedback_correct + ' ' + (question.answer_explanation || ''); feedback.className = 'feedback good'; }}
  else {{ feedback.textContent = question.feedback_retry; feedback.className = 'feedback retry'; return; }}
  if (!card.dataset.done) {{ card.dataset.done = '1'; completed += 1; updateProgress(); }}
}}
render();
</script></body></html>'''
    (directory / "index.html").write_text(document, encoding="utf-8")
    _write_json(directory / "interactive_spec.json", spec.model_dump(mode="json"))


def _render_preview(spec: SlideDeckSpec, path: Path) -> None:
    width, height = 1200, 675
    slide_count = max(1, len(spec.slides))
    rows = []
    for index, slide in enumerate(spec.slides[:12]):
        x = 40 + (index % 3) * 385
        y = 38 + (index // 3) * 156
        title = html.escape(_truncate(slide.title, 31))
        body = "".join(f"<text x=\"{x + 18}\" y=\"{y + 62 + n * 20}\" class=\"body\">{html.escape(_truncate(element.text or element.latex or element.element_type, 38))}</text>" for n, element in enumerate(slide.elements[:4]))
        rows.append(f'<g><rect x="{x}" y="{y}" width="350" height="128" rx="7" class="card"/><text x="{x + 18}" y="{y + 38}" class="title">{title}</text>{body}</g>')
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#f8f7f3"/><text x="40" y="28" class="meta">教学 Demo 预览 · {slide_count} 页</text>{''.join(rows)}<style>.card{{fill:#fff;stroke:#d8e0df}}.title{{font:700 17px 'Microsoft YaHei',sans-serif;fill:#182334}}.body{{font:14px 'Microsoft YaHei',sans-serif;fill:#52606d}}.meta{{font:13px 'Microsoft YaHei',sans-serif;fill:#1c5e70}}</style></svg>'''
    path.write_text(svg, encoding="utf-8")


def _render_markdown_summary(
    package: LoadedTeachingContentPackage,
    plan: DemoGenerationPlan,
    path: Path,
) -> None:
    """Write an audience-facing summary from IR, with traceable video evidence."""
    context = package.ir.course_context
    title = str(context.get("course_title") or plan.title or "课程内容总结")
    duration = float(context.get("duration_seconds") or package.result.metadata.duration_seconds or 0)
    lines = [
        f"# {title}",
        "",
        f"> 内容总结｜来源视频：{context.get('source_video_name', package.result.source_video.file_name)}｜时长：{_format_seconds(duration)}",
        "> 本总结以视频中文 ASR 和带时间戳的课程关键帧为证据；教学术语做了可追溯规范化，原始识别结果保留在中间包中。",
        "",
    ]
    # 一句话结论由真实数据（教学单元主题）生成。信息不足就整段省略：
    # 既不写入与本课无关的固定文本，也不用占位符顶替。
    topics = [unit.topic.strip() for unit in package.ir.teaching_units if (unit.topic or "").strip()]
    if topics:
        topic_text = "、".join(topics)
        lines.extend(
            [
                "## 一句话结论",
                "",
                f"本课程依次讲解：{topic_text}。",
                "",
            ]
        )
    lines.extend(["## 学习目标", ""])
    lines.extend(f"- {objective}" for objective in plan.learning_objectives)
    lines.extend(["", "## 核心知识点", ""])
    for unit in package.ir.teaching_units:
        time_range = f"{_format_seconds(unit.time_range.start_seconds)}–{_format_seconds(unit.time_range.end_seconds)}"
        evidence_text = _evidence_summary(unit.evidence_refs, limit=8)
        lines.extend([f"### {unit.topic}", "", f"**视频时间**：{time_range}  ", f"**证据**：{evidence_text}", ""])
        for block in unit.content_blocks:
            if block.block_type == "image":
                asset_text = ", ".join(block.asset_refs) or "未提取到图像资源"
                lines.append(f"- [视觉证据] {block.text}（资源：{asset_text}）")
                continue
            label = {"observed": "原始观察", "corrected": "术语规范化", "inferred": "上下文归纳", "unresolved": "待复核"}.get(block.status, block.status)
            lines.append(f"- [{label}] {block.text}")
        lines.append("")

    # 判断方法 / 典型现象与安全提醒 / 易错点原先是电路课固定文本，与本课无关，
    # 且没有对应的真实数据源，因此整节删除而不是替换成别的硬编码内容。
    lines.extend(
        [
            "## 证据处理说明",
            "",
            f"- 原始证据：{len(package.result.transcript.segments)} 条 ASR 片段、{len(package.result.keyframes)} 张关键帧、{len(package.result.evidence)} 条证据记录。",
            "- 处理方式：原始 ASR 不覆盖；术语规范化与上下文归纳的内容标记为 corrected 或 inferred，并保留对应时间范围和 Evidence ID。",
            "- 生成边界：PPTX、DOCX 和 HTML 均只读取 TeachingContentPackage，不重新读取原视频。",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def _format_seconds(value: float) -> str:
    seconds = max(0, int(round(float(value))))
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def _docx_bullets(document, heading: str, values: list[str]) -> None:
    document.add_heading(heading, level=1)
    for value in values or ["暂无"]:
        document.add_paragraph(value, style="List Bullet")


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def _truncate(value: str, limit: int) -> str:
    value = re.sub(r"\s+", " ", str(value)).strip()
    return value if len(value) <= limit else value[: limit - 1].rstrip() + "…"


def _estimate_pptx_text_box_height(text: str, font_size: int, width_chars: int = 25) -> float:
    """Estimate a readable legacy-PPTX text box height without clipping."""

    normalized = re.sub(r"\s+", " ", text or "").strip()
    if not normalized:
        return 0.65
    chars_per_line = max(12, int(width_chars * 18 / max(16, font_size)))
    line_count = 0
    for paragraph in normalized.split("\n"):
        line_count += max(1, (len(paragraph) + chars_per_line - 1) // chars_per_line)
    line_height_inches = 0.32 if font_size <= 18 else 0.36
    return min(2.25, max(0.65, line_count * line_height_inches + 0.16))


def _is_package_relative(path: str) -> bool:
    normalized = path.replace("\\", "/")
    return not (normalized.startswith("/") or re.match(r"^[A-Za-z]:/", normalized))
