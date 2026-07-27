import os
from docx import Document


def parse_comprehensive_word(file_path: str) -> dict:
    if not os.path.exists(file_path):
        return {"status": "error", "message": f"文件不存在: {file_path}", "data": ""}

    try:
        doc = Document(file_path)
        content_blocks = []

        # 1. 提取文字并保留直观的物理排版层级
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name

            # 遇到标题时，自动加上【标题X】的标记，并根据级别缩进
            if style_name.startswith('Heading'):
                try:
                    level = int(style_name.split(' ')[-1])
                    # 级别越低，前面的缩进越深（2级标题缩进4个空格，以此类推）
                    indent = "    " * (level - 1)
                    content_blocks.append(f"\n{indent}【标题 {level}】{text}\n")
                except ValueError:
                    content_blocks.append(f"\n【标题】{text}\n")

            # 遇到列表项目，自动加原点并缩进
            elif 'List' in style_name:
                content_blocks.append(f"    • {text}\n")

            # 普通段落直接换行
            else:
                content_blocks.append(f"{text}\n")

        # 2. 提取表格并加上明显的边框提示
        if doc.tables:
            for table in doc.tables:
                content_blocks.append("\n[--- 表格数据开始 ---]")
                for row in table.rows:
                    row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    content_blocks.append(" | ".join(row_data))
                content_blocks.append("[--- 表格数据结束 ---]\n")

        final_text = "".join(content_blocks).strip()

        return {
            "status": "success",
            "message": "Word 解析与结构保留成功",
            "data": final_text
        }

    except Exception as e:
        return {"status": "error", "message": f"Word解析失败: {str(e)}", "data": ""}


if __name__ == "__main__":
    result = parse_comprehensive_word("test_document.docx")
    print(result["data"])