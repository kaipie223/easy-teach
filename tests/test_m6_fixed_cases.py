"""M6 fixed-case regression for local retrieval, blueprints and exports."""

import asyncio
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from pptx import Presentation

from ai.rag.retriever import RAGRetriever
from backend.config import settings
from backend.schemas import EvidenceRef
from backend.services import rag as rag_service
from backend.services.courseware import compile_plan_content
from backend.services.generator import generate_docx, generate_html, generate_pptx
from backend.services.quality import inspect_courseware


FIXED_CASES = [
    {
        "id": "tcp",
        "query": "TCP 三次握手 SYN ACK",
        "source": "tcp-case.pdf",
        "title": "理解 TCP 三次握手",
        "knowledge_titles": ["连接建立过程", "报文确认机制"],
        "brief": {
            "teaching_goal": "理解 TCP 三次握手",
            "target_audience": "计算机专业大二学生",
            "duration_minutes": 45,
            "knowledge_points": [
                {
                    "order": 1,
                    "title": "连接建立过程",
                    "difficulty": "basic",
                    "key_points": ["客户端发送 SYN", "服务器返回 SYN-ACK", "客户端确认 ACK"],
                    "examples": ["客户端与服务器建立连接"],
                    "estimated_minutes": 20,
                },
                {
                    "order": 2,
                    "title": "报文确认机制",
                    "difficulty": "intermediate",
                    "key_points": ["序列号", "确认号"],
                    "examples": [],
                    "estimated_minutes": 15,
                },
            ],
            "logic_flow": ["问题导入", "过程讲解", "例题练习", "总结"],
            "teaching_focus": "三次报文交换的时序和作用",
            "teaching_difficulties": "区分 SYN 与 ACK 的含义",
            "output_types": ["pptx", "docx", "html"],
            "interaction_ideas": "让学生根据时序补全报文",
        },
    },
    {
        "id": "colors",
        "query": "认识红黄蓝 颜色认知",
        "source": "colors-case.pdf",
        "title": "认识红、黄、蓝",
        "knowledge_titles": ["认识红色", "认识黄色", "认识蓝色"],
        "brief": {
            "teaching_goal": "认识红、黄、蓝",
            "target_audience": "低年级特需儿童",
            "duration_minutes": 40,
            "knowledge_points": [
                {
                    "order": 1,
                    "title": "认识红色",
                    "difficulty": "basic",
                    "key_points": ["辨认红色", "在生活中寻找红色"],
                    "examples": ["苹果"],
                    "estimated_minutes": 8,
                },
                {
                    "order": 2,
                    "title": "认识黄色",
                    "difficulty": "basic",
                    "key_points": ["辨认黄色", "说出一个黄色物品"],
                    "examples": ["香蕉"],
                    "estimated_minutes": 8,
                },
                {
                    "order": 3,
                    "title": "认识蓝色",
                    "difficulty": "basic",
                    "key_points": ["辨认蓝色", "在图片中找出蓝色"],
                    "examples": ["天空"],
                    "estimated_minutes": 8,
                },
            ],
            "logic_flow": ["颜色导入", "观察匹配", "分类练习", "总结"],
            "teaching_focus": "稳定辨认三种基础颜色",
            "teaching_difficulties": "在不同物品和背景中保持颜色判断",
            "output_types": ["pptx", "docx", "html"],
            "interaction_ideas": "将颜色卡片与同色物品配对",
            "style_preference": "低刺激、清晰、可操作",
        },
    },
]


class LocalCaseCollection:
    """Small Chroma-shaped fixture that keeps this regression offline."""

    def __init__(self) -> None:
        self.rows = [
            {
                "id": "evidence_tcp",
                "content": "TCP 三次握手依次使用 SYN、SYN-ACK 和 ACK 建立可靠连接。",
                "case": "tcp",
                "metadata": {
                    "source": "tcp-case.pdf",
                    "evidence_id": "evidence_tcp",
                    "document_id": "case_tcp",
                    "locator_page": 3,
                },
            },
            {
                "id": "evidence_colors",
                "content": "颜色认知活动帮助学生辨认红色、黄色和蓝色，并将颜色与物品配对。",
                "case": "colors",
                "metadata": {
                    "source": "colors-case.pdf",
                    "evidence_id": "evidence_colors",
                    "document_id": "case_colors",
                    "locator_page": 1,
                },
            },
        ]

    def get(self) -> dict:
        return {
            "ids": [row["id"] for row in self.rows],
            "documents": [row["content"] for row in self.rows],
            "metadatas": [row["metadata"] for row in self.rows],
        }

    def query(self, query_texts: list[str], n_results: int) -> dict:
        query = query_texts[0].lower()
        target = "colors" if "颜色" in query or "红" in query else "tcp"
        ranked = sorted(self.rows, key=lambda row: row["case"] != target)
        return {
            "ids": [[row["id"] for row in ranked[:n_results]]],
            "documents": [[row["content"] for row in ranked[:n_results]]],
            "metadatas": [[row["metadata"] for row in ranked[:n_results]]],
            "distances": [[0.05 if row["case"] == target else 0.95 for row in ranked[:n_results]]],
        }


def local_case_retriever() -> RAGRetriever:
    retriever = RAGRetriever.__new__(RAGRetriever)
    retriever.vector_weight = 0.7
    retriever.kw_weight = 0.3
    retriever.collection = LocalCaseCollection()
    retriever._all_docs = retriever._load_all_docs()
    return retriever


def build_case_plan(case: dict, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(rag_service, "_get_retriever", local_case_retriever)
    documents = asyncio.run(rag_service.search(case["query"], top_k=2))
    assert documents
    assert documents[0].source == case["source"]
    assert documents[0].locator["page"] in {1, 3}

    refs = [
        EvidenceRef(
            evidence_id=document.evidence_id,
            document_id=document.document_id,
            source_type="local_rag",
            source_name=document.source,
            locator=document.locator,
            quote=document.content,
            score=document.score,
        )
        for document in documents
    ]
    brief = SimpleNamespace(content_json=case["brief"])
    plan = compile_plan_content(brief, refs)
    return plan, documents


@pytest.mark.parametrize("case", FIXED_CASES, ids=[case["id"] for case in FIXED_CASES])
def test_fixed_case_local_rag_and_blueprint(case, monkeypatch):
    plan, documents = build_case_plan(case, monkeypatch)

    assert plan.title == case["title"]
    assert [item.title for item in plan.knowledge_points] == case["knowledge_titles"]
    assert plan.slides[0].evidence_refs
    assert all(slide.evidence_refs for slide in plan.slides)
    assert any(document.source == case["source"] for document in documents)

    report = inspect_courseware(plan)
    assert report["status"] == "passed", report
    assert report["metrics"]["evidence_ref_count"] == len(plan.evidence_refs)


@pytest.mark.parametrize("case", FIXED_CASES, ids=[case["id"] for case in FIXED_CASES])
def test_fixed_case_exports_contain_visible_content(tmp_path, monkeypatch, case):
    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan, _ = build_case_plan(case, monkeypatch)
    payload = plan.model_dump(mode="json")

    pptx_path = Path(asyncio.run(generate_pptx(payload, output_name=f"{case['id']}.pptx")))
    docx_path = Path(asyncio.run(generate_docx(payload, output_name=f"{case['id']}.docx")))
    html_path = Path(asyncio.run(generate_html(payload, output_name=f"{case['id']}.html")))

    presentation = Presentation(str(pptx_path))
    visible_slide_text = "\n".join(
        shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")
    )
    assert case["title"] in visible_slide_text
    assert "来源：" in visible_slide_text

    document_text = "\n".join(paragraph.text for paragraph in Document(str(docx_path)).paragraphs)
    assert case["title"] in document_text
    assert "教学过程" in document_text
    assert "来源" in document_text

    html_text = html_path.read_text(encoding="utf-8")
    assert case["title"] in html_text
    assert "互动练习" in html_text
