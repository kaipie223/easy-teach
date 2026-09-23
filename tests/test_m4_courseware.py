import zipfile
from pathlib import Path

import fitz
from docx import Document

from backend.db.database import get_db
from backend.main import app
from backend.models.material import EvidenceChunk, Material
from backend.services.generator import generate_docx, generate_html, generate_pdf, generate_pptx


def register(client, email: str):
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "password123", "display_name": "M4 Teacher"},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    return payload, {"Authorization": f"Bearer {payload['access_token']}"}


def complete_brief():
    return {
        "teaching_goal": "帮助学生理解 TCP 三次握手并解释每次报文的作用",
        "target_audience": "大一新生",
        "duration_minutes": 45,
        "knowledge_points": [
            {
                "order": 1,
                "title": "连接建立过程",
                "difficulty": "basic",
                "key_points": ["SYN", "SYN-ACK", "ACK"],
                "examples": ["客户端与服务器建立连接"],
                "estimated_minutes": 20,
            },
            {
                "order": 2,
                "title": "报文确认机制",
                "difficulty": "intermediate",
                "key_points": ["序列号", "确认号"],
                "examples": [],
                "estimated_minutes": 15,
            },
        ],
        "logic_flow": ["问题导入", "过程讲解", "例题练习", "总结"],
        "teaching_focus": "三次报文交换的时序和作用",
        "teaching_difficulties": "区分 SYN 与 ACK 的含义",
        "output_types": ["pptx", "docx", "html"],
        "interaction_ideas": "让学生根据时序图补全报文",
    }


def test_courseware_plan_is_versioned_and_evidence_backed(client, monkeypatch):
    def no_external_rag(*args, **kwargs):
        return []

    monkeypatch.setattr("backend.routers.courseware.search_sync", no_external_rag)
    payload, headers = register(client, "m4-plan@example.com")
    project = client.post(
        "/api/v1/projects",
        headers=headers,
        json={"title": "M4 蓝图测试", "scenario": "通用课程"},
    )
    assert project.status_code == 201
    project_id = project.json()["project_id"]
    session = client.post(
        "/api/v1/sessions",
        headers=headers,
        json={"subject": "TCP", "project_id": project_id},
    )
    assert session.status_code == 201

    db = next(app.dependency_overrides[get_db]())
    try:
        material = Material(
            material_id="mat_m4test",
            owner_id=payload["user"]["user_id"],
            project_id=project_id,
            original_name="tcp.pdf",
            file_type="pdf",
            mime_type="application/pdf",
            stored_path="/tmp/tcp.pdf",
            size_bytes=10,
            checksum_sha256="a" * 64,
            status="ready",
            metadata_json={},
        )
        db.add(material)
        db.add(
            EvidenceChunk(
                evidence_id="evidence_m4test",
                material_id=material.material_id,
                source_type="uploaded_pdf",
                chunk_index=0,
                locator_json={"page": 3},
                text="SYN、SYN-ACK 和 ACK 完成连接建立。",
                metadata_json={},
                usage_tags=[],
                content_hash="b" * 64,
                is_valid=True,
            )
        )
        db.commit()
    finally:
        db.close()

    updated = client.patch(
        f"/api/v1/projects/{project_id}/brief",
        headers=headers,
        json=complete_brief(),
    )
    assert updated.status_code == 200, updated.text
    confirmed = client.post(
        f"/api/v1/projects/{project_id}/brief/confirm",
        headers=headers,
        json={"expected_version": updated.json()["version"]},
    )
    assert confirmed.status_code == 200, confirmed.text

    created = client.post(
        f"/api/v1/projects/{project_id}/plan",
        headers=headers,
        json={"generation_mode": "template"},
    )
    assert created.status_code == 201, created.text
    plan = created.json()
    assert plan["status"] == "ready"
    assert plan["generation_mode"] == "template"
    assert plan["version"] == 1
    assert len(plan["content"]["slides"]) >= 4
    assert plan["content"]["slides"][0]["slide_id"] == "slide_001"
    assert plan["content"]["lesson_sections"]
    assert plan["content"]["interactions"][0]["items"]
    assert plan["source_refs"][0]["evidence_id"] == "evidence_m4test"
    assert plan["source_refs"][0]["locator"]["page"] == 3

    fetched = client.get(f"/api/v1/projects/{project_id}/plan", headers=headers)
    assert fetched.status_code == 200
    assert fetched.json()["plan_id"] == plan["plan_id"]

    rebuilt = client.post(
        f"/api/v1/projects/{project_id}/plan",
        headers=headers,
        json={"force_rebuild": True, "generation_mode": "template"},
    )
    assert rebuilt.status_code == 201, rebuilt.text
    assert rebuilt.json()["version"] == 2
    assert rebuilt.json()["brief_id"] == plan["brief_id"]

    revised_content = rebuilt.json()["content"]
    revised_content["title"] = "TCP 连接建立与确认机制"
    revised_content["slides"][0]["title"] = "从一次连接请求开始"
    revised_content["output_specs"]["docx"]["homework"] = "绘制报文时序图并解释三个报文。"
    revised = client.post(
        f"/api/v1/projects/{project_id}/plan/revisions",
        headers=headers,
        json={
            "base_plan_id": rebuilt.json()["plan_id"],
            "content": revised_content,
            "summary": "调整导入页与课后任务",
        },
    )
    assert revised.status_code == 201, revised.text
    assert revised.json()["version"] == 3
    assert revised.json()["generation_mode"] == "manual"
    assert revised.json()["content"]["slides"][0]["slide_id"] == "slide_001"
    assert revised.json()["content"]["output_specs"]["docx"]["homework"].startswith("绘制")

    stale_revision = client.post(
        f"/api/v1/projects/{project_id}/plan/revisions",
        headers=headers,
        json={"base_plan_id": rebuilt.json()["plan_id"], "content": revised_content},
    )
    assert stale_revision.status_code == 409
    assert stale_revision.json()["error"]["code"] == "PLAN_VERSION_CONFLICT"


def test_four_renderers_consume_the_same_plan(tmp_path, monkeypatch):
    from backend.config import settings

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan = {
        "title": "TCP 三次握手",
        "teaching_goal": "理解连接建立过程",
        "target_audience": "大一新生",
        "duration_minutes": 45,
        "teaching_focus": "报文顺序",
        "teaching_difficulties": "SYN 与 ACK 的区别",
        "slides": [
            {
                "slide_id": "slide_001",
                "order": 1,
                "title": "TCP 三次握手",
                "purpose": "导入",
                "bullets": ["SYN", "SYN-ACK", "ACK"],
                "speaker_notes": "引入问题并观察学生已有认识",
                "evidence_refs": [{"source_name": "tcp.pdf", "locator": {"page": 3}}],
            }
        ],
        "lesson_sections": [
            {
                "section_id": "section_001",
                "order": 1,
                "title": "问题导入",
                "duration_minutes": 10,
                "objective": "提出问题",
                "teacher_actions": ["展示连接场景"],
                "student_actions": ["提出猜想"],
                "assessment": "能够复述问题",
                "evidence_refs": [],
            }
        ],
        "interactions": [
            {
                "interaction_id": "interaction_001",
                "interaction_type": "classification",
                "title": "报文排序",
                "prompt": "请选出报文",
                "items": ["SYN", "ACK", "</script><script>alert(1)</script>"],
                "answer_groups": {"核心": ["SYN", "ACK"]},
                "evidence_refs": [],
            }
        ],
        "evidence_refs": [],
        "output_specs": {
            "pptx": {
                "narrative_arc": ["问题", "解释", "练习"],
                "visual_direction": "使用时序图呈现报文方向",
                "max_bullets_per_slide": 5,
                "speaker_notes_required": True,
            },
            "docx": {
                "teacher_preparation": ["准备抓包截图"],
                "differentiation": ["为初学者提供报文提示"],
                "homework": "绘制三次握手时序图。",
                "reflection_prompts": ["学生最容易混淆哪个确认号？"],
            },
            "pdf": {
                "printable_summary": "三次握手通过 SYN、SYN-ACK 和 ACK 建立连接。",
                "assessment_checklist": ["能按顺序写出三个报文"],
                "include_sources": True,
            },
            "html": {
                "interaction_ids": ["interaction_001"],
                "completion_message": "你已经完成连接建立练习。",
                "allow_retry": False,
                "accessibility_notes": ["支持键盘操作"],
            },
        },
    }
    pptx_path = Path(generate_pptx(plan))
    docx_path = Path(generate_docx(plan))
    pdf_path = Path(generate_pdf(plan))
    html_path = Path(generate_html(plan))
    assert pptx_path.is_file() and pptx_path.stat().st_size > 0
    assert docx_path.is_file() and docx_path.stat().st_size > 0
    assert pdf_path.read_bytes().startswith(b"%PDF-")
    with fitz.open(pdf_path) as pdf:
        assert pdf.page_count == 3
        pdf_text = "".join(page.get_text() for page in pdf)
        assert "学习评价清单" in pdf_text
        assert "三次握手通过" in pdf_text
    document_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert "准备抓包截图" in document_text
    assert "绘制三次握手时序图" in document_text
    with zipfile.ZipFile(pptx_path) as archive:
        notes_xml = b"".join(
            archive.read(name) for name in archive.namelist() if name.startswith("ppt/notesSlides/")
        )
        assert "引入问题并观察学生已有认识".encode() in notes_xml
    html_text = html_path.read_text(encoding="utf-8")
    assert "报文排序" in html_text
    assert "提交答案" in html_text and "重新作答" not in html_text
    assert "你已经完成连接建立练习" in html_text
    assert "answerGroups" in html_text
    assert "</script><script>alert(1)</script>" not in html_text


# ── 幻灯片配图 ────────────────────────────────────────


def _make_png(path, size=(1200, 400)):
    """A deliberately wide picture so cover-cropping has to do real work."""
    from PIL import Image

    Image.new("RGB", size, (18, 99, 255)).save(path)
    return path


def _slide_with_image(
    placement: str,
    *,
    caption: str = "",
    material_id: str = "mat_pic",
    slide_id: str = "slide_001",
    order: int = 1,
) -> dict:
    return {
        "slide_id": slide_id,
        "order": order,
        "title": "配图页",
        "purpose": "演示配图",
        "bullets": ["第一条要点", "第二条要点"],
        "speaker_notes": "讲稿正文",
        "evidence_refs": [],
        "image": {"material_id": material_id, "placement": placement, "caption": caption},
    }


def _plan_with_image(*slides: dict) -> dict:
    return {
        "title": "配图课程",
        "teaching_goal": "验证配图渲染",
        "target_audience": "大一新生",
        "duration_minutes": 45,
        "slides": list(slides),
        "lesson_sections": [],
        "interactions": [],
        "evidence_refs": [],
        "output_specs": {"pptx": {"max_bullets_per_slide": 5, "speaker_notes_required": True}},
    }


def _slide_text(slide) -> str:
    return "".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def _cover_slide(*, slide_id: str = "slide_001", order: int = 1) -> dict:
    """一页没有配图的封面，让测试牌组贴近真实生成结果。"""
    return {
        "slide_id": slide_id,
        "order": order,
        "title": "封面",
        "purpose": "课程导入",
        "layout": "cover",
        "bullets": [],
        "speaker_notes": "封面讲稿",
        "evidence_refs": [],
    }


def _deck_with_picture(placement: str, *, caption: str = "") -> dict:
    """封面 + 一页带配图的讲解页。"""
    return _plan_with_image(
        _cover_slide(),
        _slide_with_image(placement, caption=caption, slide_id="slide_002", order=2),
    )


def _full_bleed_shapes(slide, presentation) -> list:
    """整页大小的形状，用来把"背景蒙层"和普通装饰块区分开。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and shape.width == presentation.slide_width
        and shape.height == presentation.slide_height
    ]


def test_pptx_embeds_each_slide_picture_placement(tmp_path, monkeypatch):
    """三种位置模式都要真正嵌入图片，并各自呈现正确的版面行为。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    image = _make_png(tmp_path / "figure.png")

    for placement in ("right", "full", "background"):
        plan = _deck_with_picture(placement, caption="示意图")
        path = Path(
            generate_pptx(plan, output_name=f"{placement}.pptx", images={"mat_pic": image})
        )
        presentation = Presentation(path)
        # 第 1 页是封面，第 2 页才承载配图
        assert not [
            s for s in presentation.slides[0].shapes
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        slide = presentation.slides[1]
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pictures) == 1, placement
        picture = pictures[0]
        body_text = _slide_text(slide)
        notes_text = slide.notes_slide.notes_text_frame.text
        full_bleed = _full_bleed_shapes(slide, presentation)

        if placement in {"background", "full"}:
            # 满版铺满只能靠裁剪把宽图收成 16:9，否则会被拉伸变形
            assert picture.crop_left > 0 and picture.crop_right > 0
            assert picture.width == presentation.slide_width
            assert picture.height == presentation.slide_height
            # 满版图上必须叠一层整页遮罩，否则文字读不出来
            assert len(full_bleed) == 1, placement
            # 而且必须是"白→透明"渐变：纯色遮罩会把任何图片压成一片灰白
            scrim = full_bleed[0]._element.spPr
            gradient = scrim.find(qn("a:gradFill"))
            assert gradient is not None, "满版遮罩必须是渐变蒙层"
            assert len(gradient.find(qn("a:gsLst")).findall(qn("a:gs"))) == 3, "渐变要有多个停点"
            # 图注满版放不下，改为并入讲稿
            assert "示意图" not in body_text
            assert "示意图" in notes_text
        else:
            # 等比缩放，不能靠裁剪
            assert picture.crop_left == 0 and picture.crop_right == 0
            assert picture.width < presentation.slide_width
            assert not full_bleed, placement
            assert "示意图" in body_text

        # 满版页同样把要点留在画面上：文字压在渐变蒙层上，不再被赶进讲稿
        assert "第一条要点" in body_text, placement
        assert "第一条要点" not in notes_text, placement

        assert "讲稿正文" in notes_text

    # 右侧图文要给图片让出正文宽度；两种满版模式（background / full）下要点占满整宽。
    def bullet_box_width(placement: str) -> int:
        slide = Presentation(
            Path(
                generate_pptx(
                    _deck_with_picture(placement),
                    output_name=f"{placement}-width.pptx",
                    images={"mat_pic": image},
                )
            )
        ).slides[1]
        boxes = [
            shape
            for shape in slide.shapes
            if shape.has_text_frame and "第一条要点" in shape.text_frame.text
        ]
        assert boxes, placement
        return boxes[0].width

    assert bullet_box_width("right") < bullet_box_width("background")
    assert bullet_box_width("background") == bullet_box_width("full")


def test_pptx_renders_without_a_picture_when_it_cannot_be_used(tmp_path, monkeypatch):
    """素材被删除或文件丢失时导出必须照常成功：跳过配图，要点回到常规版式。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan = _deck_with_picture("full")

    # 没有解析到图片（素材已归档 / 不属于本项目）
    nothing = Presentation(
        Path(generate_pptx(plan, output_name="none.pptx", images={}))
    ).slides[1]
    # 拿到了路径但文件已经不在磁盘上
    missing = Presentation(
        Path(
            generate_pptx(
                plan,
                output_name="gone.pptx",
                images={"mat_pic": tmp_path / "gone.png"},
            )
        )
    ).slides[1]

    for slide in (nothing, missing):
        assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert "第一条要点" in _slide_text(slide)
        assert "讲稿正文" in slide.notes_slide.notes_text_frame.text


def test_pptx_applies_the_visual_system(tmp_path, monkeypatch):
    """版式系统：CJK 字体、显式配色、按 layout 区分版式、不再有"暂无证据"噪音。"""
    import zipfile

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from backend.config import settings
    from backend.services.generator import COLOR_BAND, COLOR_INK, generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    summary = _slide_with_image("right", slide_id="slide_003", order=3)
    summary.pop("image")
    summary["layout"] = "summary"
    summary["evidence_refs"] = [
        {"source_type": "pdf", "source_name": "教材.pdf", "locator": {"page": 3}}
    ]
    plan = _plan_with_image(_cover_slide(), _slide_with_image("right", slide_id="slide_002", order=2), summary)

    path = Path(generate_pptx(plan, output_name="design.pptx"))

    with zipfile.ZipFile(path) as archive:
        cover = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        # 中文字体必须连东方字体一起设置，否则 PowerPoint 退回主题字体
        assert "<a:ea" in cover and 'typeface="Microsoft YaHei"' in cover
        assert "<a:cs" in cover
        # 主色必须显式写入，不能依赖主题色
        assert "1463FF" in cover.upper()
        # 没有证据的页面不许出现「暂无可用证据」
        for name in ("ppt/slides/slide1.xml", "ppt/slides/slide2.xml"):
            assert "暂无可用证据" not in archive.read(name).decode("utf-8")
        # 有证据的页面照常显示来源
        assert "教材.pdf" in archive.read("ppt/slides/slide3.xml").decode("utf-8")

    presentation = Presentation(path)
    cover_slide, body_slide, summary_slide = presentation.slides
    cover_blocks = [s for s in cover_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    # 封面有一条贯穿整页的深色带
    assert len(_full_bleed_shapes(cover_slide, presentation)) == 0
    assert any(shape.height == presentation.slide_height for shape in cover_blocks)
    # 标准页有顶部品牌条 + 标题下划线，且没有整页色块
    body_kinds = [s.shape_type for s in body_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert body_kinds
    assert not _full_bleed_shapes(body_slide, presentation)
    # 小结页用深色带色的标题，标准讲解页用正文色：两者是可区分的版式
    def title_color(slide) -> str:
        return str(slide.shapes.title.text_frame.paragraphs[0].font.color.rgb)

    assert title_color(summary_slide) == COLOR_BAND
    assert title_color(body_slide) == COLOR_INK


def _labelled_slide(layout: str, *, order: int, bullets=None, notes: str = "讲稿正文") -> dict:
    slide = _slide_with_image("right", slide_id=f"slide_{order:03d}", order=order)
    slide["layout"] = layout
    slide["speaker_notes"] = notes
    if bullets is not None:
        slide["bullets"] = bullets
    return slide


def test_pptx_uses_named_layouts_placeholders_and_theme(tmp_path, monkeypatch):
    """改用命名版式 + 占位符 + 主题，而不是在空白版式上手画文本框。"""
    import zipfile

    from pptx import Presentation

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    deck = _plan_with_image(
        _cover_slide(),
        _labelled_slide("bullets", order=2),
        _labelled_slide("agenda", order=3),
        _labelled_slide("summary", order=4),
    )
    path = Path(generate_pptx(deck, output_name="layouts.pptx"))
    presentation = Presentation(path)

    # 每一页都从命名版式创建，不再一律用 Blank
    # （注意：Slides 的切片返回的不是 Slide 对象，所以逐页取下标）
    layout_names = [presentation.slides[i].slide_layout.name for i in range(4)]
    assert layout_names == ["Title Slide", "Title and Content", "Title and Content", "Title and Content"]

    # 标题落在版式自带的占位符里，PowerPoint 的大纲视图与"重设版式"才有意义
    for slide in presentation.slides:
        title = slide.shapes.title
        assert title is not None
        assert title.is_placeholder
        assert title.text_frame.text.strip()

    with zipfile.ZipFile(path) as archive:
        theme = archive.read("ppt/theme/theme1.xml").decode("utf-8")
        # 主题配色与中文字体都要落进 theme1.xml，占位符和表格才会自动跟随
        assert "1463FF" in theme.upper()
        assert theme.count("Microsoft YaHei") >= 6
        # script="Hans" 的优先级高于 a:ea，中文必须两处都改
        assert 'script="Hans" typeface="Microsoft YaHei"' in theme


def test_pptx_renders_keyword_emphasis(tmp_path, monkeypatch):
    """要点里的 **关键词** 变成加粗变色，且标记不会泄漏成裸星号。"""
    from pptx import Presentation

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan = _plan_with_image(
        _cover_slide(),
        _labelled_slide(
            "bullets",
            order=2,
            bullets=["**化合价升降**是电子转移的外显结果", "普通要点"],
        ),
    )

    slide = Presentation(Path(generate_pptx(plan, output_name="emphasis.pptx"))).slides[1]

    emphasised = [
        run
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.strip() == "化合价升降"
    ]
    assert len(emphasised) == 1
    assert emphasised[0].font.bold is True
    assert str(emphasised[0].font.color.rgb) == "1463FF"

    rendered = "".join(
        run.text
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    )
    assert "**" not in rendered


def test_pptx_layout_skeleton_applies_without_model_labels(tmp_path, monkeypatch):
    """模型没标 layout 时也必须出现版式差异，不能 N 页长得一样。"""
    from pptx import Presentation

    from backend.config import settings
    from backend.services.generator import COLOR_BAND, COLOR_INK, generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    # 刻意不给 layout：这正是旧提示词下真实蓝图的样子
    plan = _plan_with_image(
        _cover_slide(),
        _slide_with_image("right", slide_id="slide_002", order=2),
        _slide_with_image("right", slide_id="slide_003", order=3),
        _slide_with_image("right", slide_id="slide_004", order=4),
    )

    presentation = Presentation(Path(generate_pptx(plan, output_name="skeleton.pptx")))

    def title_color(slide) -> str:
        return str(slide.shapes.title.text_frame.paragraphs[0].font.color.rgb)

    assert presentation.slides[0].slide_layout.name == "Title Slide"
    # 末页被确定性骨架改成了小结，中间页仍是标准讲解页
    assert title_color(presentation.slides[3]) == COLOR_BAND
    assert title_color(presentation.slides[1]) == COLOR_INK


def test_pptx_renders_step_and_compare_layouts(tmp_path, monkeypatch):
    """编号步骤与双栏对比必须真的按各自版式渲染，而不是退回圆点列表。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan = _plan_with_image(
        _cover_slide(),
        _labelled_slide("steps", order=2, bullets=["先读结构", "再写前序", "最后写中序"]),
        _labelled_slide("compare", order=3, bullets=["左一", "左二", "右一", "右二"]),
    )

    presentation = Presentation(Path(generate_pptx(plan, output_name="layouts2.pptx")))

    steps = presentation.slides[1]
    badges = [
        shape
        for shape in steps.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and shape.has_text_frame
        and shape.text_frame.text.strip() in {"1", "2", "3"}
    ]
    assert len(badges) == 3
    # 导语来自 purpose，直接投影给学生看
    assert "演示配图" in _slide_text(steps)

    compare = presentation.slides[2]
    panels = [
        shape
        for shape in compare.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.width == Inches(5.72)
    ]
    assert len(panels) == 2
    compare_text = _slide_text(compare)
    for item in ("左一", "左二", "右一", "右二"):
        assert item in compare_text


def test_resolve_slide_images_only_returns_usable_pictures_of_the_project(
    client, tmp_path
):
    """只解析属于本项目、未归档、确实是图片且文件仍在的引用。"""
    from datetime import datetime, timezone

    from backend.services.materials import resolve_slide_images

    payload, headers = register(client, "m4-slide-images@example.com")
    owner_id = payload["user"]["user_id"]
    project_id = client.post(
        "/api/v1/projects",
        headers=headers,
        json={"title": "配图归属测试", "scenario": "通用课程"},
    ).json()["project_id"]
    other_project_id = client.post(
        "/api/v1/projects",
        headers=headers,
        json={"title": "另一个项目", "scenario": "通用课程"},
    ).json()["project_id"]

    keep = _make_png(tmp_path / "keep.png")
    archived = _make_png(tmp_path / "archived.png")
    foreign = _make_png(tmp_path / "foreign.png")
    document = tmp_path / "notes.pdf"
    document.write_bytes(b"%PDF-1.4\n")

    db = next(app.dependency_overrides[get_db]())
    try:
        def material(material_id, project, name, file_type, stored_path, **extra):
            return Material(
                material_id=material_id,
                owner_id=owner_id,
                project_id=project,
                original_name=name,
                file_type=file_type,
                stored_path=str(stored_path),
                size_bytes=1,
                checksum_sha256=material_id,
                **extra,
            )

        db.add_all(
            [
                material("mat_keep", project_id, "keep.png", "image", keep),
                material("mat_gone", project_id, "gone.png", "image", tmp_path / "gone.png"),
                material(
                    "mat_archived",
                    project_id,
                    "archived.png",
                    "image",
                    archived,
                    deleted_at=datetime.now(timezone.utc),
                ),
                material("mat_foreign", other_project_id, "foreign.png", "image", foreign),
                material("mat_pdf", project_id, "notes.pdf", "pdf", document),
            ]
        )
        db.commit()

        plan = _plan_with_image(
            _slide_with_image("right", material_id="mat_keep"),
            _slide_with_image("full", material_id="mat_gone", slide_id="slide_002", order=2),
            _slide_with_image("full", material_id="mat_archived", slide_id="slide_003", order=3),
            _slide_with_image("full", material_id="mat_foreign", slide_id="slide_004", order=4),
            _slide_with_image("full", material_id="mat_pdf", slide_id="slide_005", order=5),
        )

        resolved = resolve_slide_images(db, project_id, plan)

        assert set(resolved) == {"mat_keep"}
        assert resolved["mat_keep"] == keep
    finally:
        db.close()


def test_resolve_slide_images_tolerates_a_non_plan_snapshot(db_session_factory):
    """旧版匿名会话传的是生成指令字典，这里必须表示"没有图片"而不是报错。"""
    from backend.services.materials import resolve_slide_images

    legacy = {"teaching_goal": "旧格式", "target_audience": "大一新生", "duration_minutes": 45}
    db = db_session_factory()
    try:
        assert resolve_slide_images(db, "p_anything", legacy) == {}
    finally:
        db.close()


def test_lesson_document_and_handout_carry_a_slide_picture_appendix(tmp_path, monkeypatch):
    """教案与打印版都要把幻灯片配图作为附录带上；没有配图时不加这一章。"""
    from backend.config import settings
    from backend.services.generator import generate_docx, generate_pdf

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    image = _make_png(tmp_path / "figure.png")
    plan = _plan_with_image(
        _slide_with_image("right", caption="两次读数对比"),
        _slide_with_image("full", slide_id="slide_002", order=2),
    )
    images = {"mat_pic": image}

    docx_path = Path(generate_docx(plan, output_name="with-image.docx", images=images))
    text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert "八、幻灯片配图" in text
    assert "两次读数对比" in text
    assert "第 2 页 · 配图页" in text
    with zipfile.ZipFile(docx_path) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())

    pdf_path = Path(generate_pdf(plan, output_name="with-image.pdf", images=images))
    with fitz.open(pdf_path) as pdf:
        # 1 页正文 + 2 页配图附录
        assert pdf.page_count == 3
        assert pdf[0].get_images() == []
        assert pdf[1].get_images() and pdf[2].get_images()

    # 解析不到配图时不产生附录章节，既有页数也不变
    plain_docx = Path(generate_docx(plan, output_name="plain.docx"))
    plain_text = "\n".join(paragraph.text for paragraph in Document(plain_docx).paragraphs)
    assert "八、幻灯片配图" not in plain_text
    with fitz.open(generate_pdf(plan, output_name="plain.pdf")) as pdf:
        assert pdf.page_count == 1


# ── 结构化版式 ────────────────────────────────────────


def _layout_plan(*slides: dict) -> dict:
    return {
        "title": "版式课程",
        "teaching_goal": "验证结构化版式",
        "target_audience": "大一新生",
        "duration_minutes": 45,
        "slides": list(slides),
        "lesson_sections": [],
        "interactions": [],
        "evidence_refs": [],
        "output_specs": {"pptx": {"max_bullets_per_slide": 6, "speaker_notes_required": True}},
    }


def _layout_slide(layout: str, bullets: list) -> dict:
    return {
        "slide_id": "slide_001",
        "order": 1,
        "title": "版式页",
        "purpose": "导语",
        "layout": layout,
        "bullets": bullets,
        "speaker_notes": "讲稿正文",
        "evidence_refs": [],
    }


def _render_one_slide(tmp_path, monkeypatch, layout: str, bullets: list):
    """渲染单页课件并返回该页。单页时不会触发封面／小结的位置兜底。"""
    from pptx import Presentation

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    path = generate_pptx(_layout_plan(_layout_slide(layout, bullets)), output_name=f"{layout}.pptx")
    return Presentation(Path(path)).slides[0]


def _auto_shapes(slide) -> list:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]


def _font_size_of(slide, target: str) -> float:
    """读某个文本的显式字号。

    `_add_textbox` 把字号写在段落级 `defRPr`，`_style_run` 写在 run 级，所以要
    两级都看——这也正是 PowerPoint 的取值顺序（run 覆盖段落）。
    """
    found = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip() == target
    ]
    assert found, f"页面上找不到「{target}」"
    paragraph = found[0].text_frame.paragraphs[0]
    run_sizes = [run.font.size for run in paragraph.runs if run.font.size is not None]
    size = run_sizes[0] if run_sizes else paragraph.font.size
    assert size is not None, f"「{target}」没有显式字号"
    return size.pt


def test_cards_layout_gives_every_item_its_own_panel(tmp_path, monkeypatch):
    """并列卡片：每条要点一张面板，冒号前半作小标题、后半作说明。"""
    slide = _render_one_slide(
        tmp_path,
        monkeypatch,
        "cards",
        ["热带雨林：全年高温多雨", "温带季风：夏季高温多雨", "地中海：夏季炎热干燥"],
    )
    # 3 张卡片 + 3 条强调线 + 品牌条 + 标题强调线
    assert len(_auto_shapes(slide)) == 8
    text = _slide_text(slide)
    assert "热带雨林" in text and "全年高温多雨" in text
    assert "地中海" in text


def test_flow_layout_connects_stages_with_arrows(tmp_path, monkeypatch):
    """流程：3 段内容画成 3 个块，块间的箭头是与步骤页唯一的可见差别。"""
    slide = _render_one_slide(
        tmp_path, monkeypatch, "flow", ["铁与氧气接触", "生成氧化铁", "结构疏松剥落"]
    )
    # 3 个流程块 + 2 个箭头 + 品牌条 + 标题强调线
    assert len(_auto_shapes(slide)) == 7
    assert "生成氧化铁" in _slide_text(slide)


def test_metric_layout_enlarges_the_leading_number(tmp_path, monkeypatch):
    """度量：开头的数字被放大；提不出数字的条目退回普通短句，内容不能丢。"""
    slide = _render_one_slide(
        tmp_path, monkeypatch, "metric", ["70%：能在 5 分钟内完成", "本题共 12 分"]
    )
    assert _font_size_of(slide, "70%") == 44
    text = _slide_text(slide)
    # 没有前导数字的那条整条按普通短句排，既没消失也没被当成数字
    assert "本题共 12 分" in text
    assert "能在 5 分钟内完成" in text


def test_quote_layout_keeps_attribution_separate(tmp_path, monkeypatch):
    """引用：出处单独小字排在下方，不与引文正文同号同重。"""
    slide = _render_one_slide(
        tmp_path, monkeypatch, "quote", ["学而时习之，不亦说乎。", "—— 《论语·学而》"]
    )
    # 引文竖条 + 品牌条 + 标题强调线
    assert len(_auto_shapes(slide)) == 3
    text = _slide_text(slide)
    assert "学而时习之" in text
    assert "— 《论语·学而》" in text
    assert _font_size_of(slide, "学而时习之，不亦说乎。") > _font_size_of(slide, "— 《论语·学而》")


def test_layout_overflow_is_carried_into_the_speaker_notes(tmp_path, monkeypatch):
    """版面装不下的条目必须并入讲稿：版式容量不该让内容凭空消失。"""
    from pptx import Presentation

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan = _layout_plan(_layout_slide("cards", [f"第 {index} 项：说明" for index in range(1, 6)]))
    slide = Presentation(Path(generate_pptx(plan, output_name="overflow.pptx"))).slides[0]

    # 卡片上限 4 张（4 张卡片 + 4 条强调线 + 品牌条 + 标题线），第 5 条不上画面
    assert len(_auto_shapes(slide)) == 10
    assert "第 5 项" not in _slide_text(slide)
    notes = slide.notes_slide.notes_text_frame.text
    assert "第 5 项" in notes
    assert "讲稿正文" in notes


def test_card_label_split_keeps_long_clauses_whole():
    """拆词只在冒号足够靠前时才生效，长从句不能被当成小标题。"""
    from backend.services.generator import _split_card_label

    assert _split_card_label("热带雨林：全年高温多雨") == ("热带雨林", "全年高温多雨")
    # 前半是一整个从句，拆出来会得到一个又长又不像标题的"小标题"
    clause = "光合作用的基本过程：叶绿体吸收光能并转化为化学能"
    assert _split_card_label(clause) == ("", clause)
    assert _split_card_label("只有一句话") == ("", "只有一句话")


def test_layout_names_are_normalised_before_rendering():
    """同义词在写入快照时就归一；未识别与未指定都要原样保留，以免封面兜底失效。"""
    from backend.schemas import SlideSpec, normalize_layout
    from backend.services.generator import _slide_layout

    assert normalize_layout("Flow Chart") == "flow"
    assert normalize_layout("two-column") == "compare"
    assert normalize_layout("GRID") == "cards"
    assert normalize_layout("diagram") == "diagram"
    assert SlideSpec(
        slide_id="s", order=1, title="t", purpose="p", layout="Flowchart"
    ).layout == "flow"

    # 归一之后仍由渲染器决定位置兜底
    assert _slide_layout({}, 0, 6) == "cover"
    assert _slide_layout({}, 5, 6) == "summary"
    assert _slide_layout({"layout": "diagram"}, 3, 6) == "bullets"
    assert _slide_layout({"layout": "two_column"}, 3, 6) == "compare"


def test_structural_layout_steps_aside_for_a_picture(tmp_path, monkeypatch):
    """结构化版式带配图时退回标准页：卡片被右图挤成窄条比没有版式更难看。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from backend.config import settings
    from backend.services.generator import generate_pptx

    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    image = _make_png(tmp_path / "figure.png")
    slide_spec = _slide_with_image("right")
    slide_spec["layout"] = "cards"

    slide = Presentation(
        Path(
            generate_pptx(
                _plan_with_image(slide_spec),
                output_name="cards-with-image.pptx",
                images={"mat_pic": image},
            )
        )
    ).slides[0]
    assert len([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1
    # 退回标准页后要点留在画面上，没有按卡片容量被截断
    assert "第一条要点" in _slide_text(slide)
