"""质量优先改造的契约：去上限、渲染层承接、教师要求注入。

这一组行为是配套的，单独看任何一条都容易被误改回去：
- 生成端不再限制每页条数与总页数；
- 渲染端必须承接"一页装不下"，拆成续页而不是截断；
- 翔实度只报 warning，密度不再报 error；
- 教师的明确要求要进提示词，并有一条确定性检查兜底。
"""

from pathlib import Path
from types import SimpleNamespace

import pytest
from openai import APIStatusError
from pydantic import ValidationError
from pptx import Presentation

from backend.config import settings
from backend.schemas import CoursewarePlanBuildRequest, EvidenceRef, PptxContentSpec
from backend.services import courseware_ai
from backend.services.ai_completion import json_completion
from backend.services.courseware import compile_plan_content
from backend.services.courseware_ai import (
    _completion,
    _make_completer,
    _normalize_and_validate,
    generate_courseware_spec_agentic,
)
from backend.services.generator import generate_pptx
from backend.services.prompt_library import (
    build_generation_prompt,
    build_requirement_block,
)
from backend.services.quality import (
    MIN_CHECKLIST_ITEMS,
    MIN_PRINTABLE_SUMMARY_CHARS,
    MIN_SPEAKER_NOTES_CHARS,
    inspect_courseware,
)

BRIEF_CONTENT = {
    "teaching_goal": "理解浮力并解释生活现象",
    "subject": "物理",
    "grade": "初二",
    "target_audience": "初二学生",
    "duration_minutes": 45,
    "knowledge_points": [
        {"order": 1, "title": "浮力概念", "estimated_minutes": 20},
        {"order": 2, "title": "阿基米德原理", "estimated_minutes": 25},
    ],
    "logic_flow": ["导入", "新授", "小结"],
    "teaching_focus": "影响浮力大小的因素",
    "teaching_difficulties": "排开液体体积的理解",
    "output_types": ["pptx", "docx", "pdf", "html"],
}


def _payload(*, slide_count: int = 6, bullets_per_slide: int = 3) -> dict:
    return {
        "title": "浮力",
        "slides": [
            {
                "order": index,
                "layout": "cover" if index == 1 else "bullets",
                "title": f"第 {index} 页",
                "purpose": f"第 {index} 页的导语",
                "bullets": [
                    f"第 {index} 页第 {position} 条需要完整投影的要点内容"
                    for position in range(1, bullets_per_slide + 1)
                ],
                "speaker_notes": "讲稿内容。" * 20,
            }
            for index in range(1, slide_count + 1)
        ],
        "lesson_sections": [
            {
                "order": 1,
                "title": "导入",
                "duration_minutes": 5,
                "objective": "学生能提出可探究的问题",
                "teacher_actions": ["提问原话"],
                "student_actions": ["写出猜想"],
                "assessment": "能写出猜想",
            },
            {
                "order": 2,
                "title": "新授",
                "duration_minutes": 35,
                "objective": "学生能说出浮力的影响因素",
                "teacher_actions": ["演示实验"],
                "student_actions": ["记录数据"],
                "assessment": "能说出控制变量",
            },
            {
                "order": 3,
                "title": "小结",
                "duration_minutes": 5,
                "objective": "学生能复述原理",
                "teacher_actions": ["组织复述"],
                "student_actions": ["复述原理"],
                "assessment": "复述完整",
            },
        ],
        "interactions": [
            {
                "interaction_type": "quiz",
                "title": "判断说法",
                "prompt": "下面哪个说法正确？",
                "items": ["说法 A", "说法 B"],
                "answer_groups": {"正确": ["说法 A"]},
            }
        ],
        "output_specs": {
            "pptx": {
                "max_bullets_per_slide": 5,
                "narrative_arc": ["导入", "新授", "小结"],
                "visual_direction": "简洁",
            },
            "docx": {
                "homework": "用一个新情境解释浮力并说明条件",
                "teacher_preparation": ["弹簧测力计", "烧杯"],
                "differentiation": ["给出分步提示"],
                "reflection_prompts": ["哪一步卡住"],
            },
            "pdf": {
                "printable_summary": "浮力的关键术语与适用条件。" * 10,
                "assessment_checklist": ["能说明浮力方向"],
            },
            "html": {
                "completion_message": "完成",
                "allow_retry": True,
                "accessibility_notes": ["键盘可操作"],
            },
        },
    }


def _template_plan() -> dict:
    brief = SimpleNamespace(content_json=BRIEF_CONTENT)
    return compile_plan_content(brief, []).model_dump(mode="json")


# ── 生成端：密度与页数不再是"裁剪线" ─────────────────────


def test_schema_no_longer_caps_page_density():
    assert PptxContentSpec(max_bullets_per_slide=20).max_bullets_per_slide == 20
    # 下限放宽到 1：只有一条要点的页是合法的，不该被判失败
    assert PptxContentSpec(max_bullets_per_slide=1).max_bullets_per_slide == 1
    with pytest.raises(ValidationError):
        PptxContentSpec(max_bullets_per_slide=0)


def test_dense_slide_passes_whole_deck_validation():
    """一页写 12 条不再判失败，声明值被抬到实际条数。"""
    spec = _normalize_and_validate(
        _payload(bullets_per_slide=12),
        brief_content=BRIEF_CONTENT,
        evidence_refs=[],
    )

    assert len(spec.slides[0].bullets) == 12
    # 声明值描述密度，不是一条能把内容砍掉的裁剪线
    assert spec.output_specs.pptx.max_bullets_per_slide == 12


def test_blueprint_page_count_is_decided_by_content():
    """页数不再被 6-16 的区间卡死：两页的课与二十页的课都合法。"""
    small = _normalize_and_validate(
        _payload(slide_count=2), brief_content=BRIEF_CONTENT, evidence_refs=[]
    )
    large = _normalize_and_validate(
        _payload(slide_count=20), brief_content=BRIEF_CONTENT, evidence_refs=[]
    )

    assert len(small.slides) == 2
    assert len(large.slides) == 20


# ── 渲染端：装不下就分页，绝不静默截断 ────────────────────


def test_renderer_splits_overfull_slides_into_continuation_pages(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan = _template_plan()
    dense = [f"要点 {index}：这是第 {index} 条需要完整投影的内容" for index in range(1, 13)]
    plan["slides"][0]["layout"] = "bullets"
    plan["slides"][0]["title"] = "密集页"
    plan["slides"][0]["bullets"] = dense
    plan["output_specs"]["pptx"]["max_bullets_per_slide"] = 5

    presentation = Presentation(str(generate_pptx(plan, output_name="dense.pptx")))
    page_texts = [
        "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        for slide in presentation.slides
    ]
    all_text = "\n".join(page_texts)

    # 12 条按每页 5 条拆成 3 页：拆是"继续投影"，不是扔掉
    assert all(bullet in all_text for bullet in dense)
    assert sum("（续 " in text for text in page_texts) == 2


def test_continuation_pages_use_the_standard_layout(tmp_path, monkeypatch):
    """续页退回标准讲解版式，且不重复携带同一张配图。"""
    monkeypatch.setattr(settings, "output_dir", Path(tmp_path))
    plan = _template_plan()
    plan["slides"][0]["layout"] = "summary"
    plan["slides"][0]["title"] = "小结页"
    plan["slides"][0]["bullets"] = [f"收束要点 {index}" for index in range(1, 8)]
    plan["output_specs"]["pptx"]["max_bullets_per_slide"] = 4

    presentation = Presentation(str(generate_pptx(plan, output_name="summary-split.pptx")))
    titles = [
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text") and "（续 " in shape.text
    ]

    assert titles == ["小结页（续 1）"]


def test_renderer_shrinks_the_body_font_for_dense_pages(tmp_path, monkeypatch):
    """密度放开后字号必须跟着内容量走，否则要点会压出页面。"""
    from backend.services.generator import _fit_body_size

    box = {"width_in": 11.3, "height_in": 4.3}
    sparse = _fit_body_size(["一句话要点。"], **box)
    dense = _fit_body_size([("很长的要点内容，包含完整的说明与判断依据。" * 6)] * 8, **box)

    assert sparse == 20.0
    assert dense < sparse
    # 再密也要留在可读范围内
    assert dense >= 12.0


# ── 质检层：密度不再是错误，翔实度会被指出来 ──────────────


def test_quality_no_longer_errors_on_dense_slides():
    plan = _template_plan()
    plan["slides"][0]["bullets"] = [f"要点 {index}" for index in range(1, 13)]

    report = inspect_courseware(plan)

    error_codes = {item["code"] for item in report["errors"]}
    assert "PPTX_TOO_MANY_BULLETS" not in error_codes
    assert report["status"] != "failed"


def test_quality_points_out_thin_content_as_warning():
    plan = _template_plan()
    plan["slides"][0]["speaker_notes"] = "讲解本页。"
    plan["output_specs"]["pdf"]["printable_summary"] = "很短。"
    plan["output_specs"]["pdf"]["assessment_checklist"] = ["只有一条"]

    report = inspect_courseware(plan)

    codes = {item["code"] for item in report["warnings"]}
    assert {"THIN_SPEAKER_NOTES", "PDF_THIN_SUMMARY", "PDF_THIN_CHECKLIST"} <= codes
    # 翔实度是质量判断，不该把一份能用的蓝图卡死
    assert report["errors"] == []


def test_forbidden_content_is_reported_against_the_brief():
    plan = _template_plan()
    plan["slides"][0]["bullets"] = ["用微积分的思想理解变化率"]

    report = inspect_courseware(plan, brief_content={"forbidden_content": "不要涉及微积分"})

    codes = {item["code"] for item in report["warnings"]}
    assert "FORBIDDEN_CONTENT_PRESENT" in codes
    # 分词是启发式的，所以只提示不判失败：误判不该让整份蓝图作废
    assert report["errors"] == []


def test_forbidden_content_check_is_skipped_without_a_brief():
    plan = _template_plan()
    plan["slides"][0]["bullets"] = ["用微积分的思想理解变化率"]

    report = inspect_courseware(plan)

    codes = {item["code"] for item in report["warnings"]}
    assert "FORBIDDEN_CONTENT_PRESENT" not in codes


# ── 教师要求注入 ────────────────────────────────────────


def test_requirement_block_lists_teacher_requirements():
    block = build_requirement_block(
        {
            "forbidden_content": "不要涉及微积分",
            "extra_requirements": "每个环节都要有板书设计",
            "style_preference": "活泼",
            "output_types": ["pptx", "docx"],
        }
    )

    assert "不要涉及微积分" in block
    assert "每个环节都要有板书设计" in block
    assert "活泼" in block
    # 列表字段要展开成可读文本，而不是 Python 列表字面量
    assert "pptx、docx" in block
    assert "[" not in block


def test_requirement_block_is_empty_without_requirements():
    assert build_requirement_block(None) == ""
    assert build_requirement_block({}) == ""
    assert build_requirement_block({"forbidden_content": "   "}) == ""


def test_generation_prompt_carries_requirements_without_dropping_generic_rules():
    plain = build_generation_prompt("物理", "初二")
    assert "禁止出现的内容" not in plain

    with_requirements = build_generation_prompt(
        "物理", "初二", {"forbidden_content": "不要涉及微积分"}
    )
    assert "微积分" in with_requirements
    # 通用规则不能被要求块挤掉
    assert "只返回 JSON 对象" in with_requirements
    # 学科与学段块同理
    assert "变量控制" in with_requirements


# ── 输出预算：给足，但被拒时要能降级 ──────────────────────


def _rejection(message: str) -> APIStatusError:
    """造一个"供应商拒绝参数"的异常。

    直接实例化 APIStatusError 需要伪造 httpx.Response（以及它的 request），
    这里只关心 message，所以用子类绕过父类的初始化。
    """

    class _Rejection(APIStatusError):
        def __init__(self, text: str) -> None:
            Exception.__init__(self, text)
            self.message = text

    return _Rejection(message)


def test_completion_lowers_the_budget_when_the_provider_rejects_it(monkeypatch):
    monkeypatch.setattr(settings, "deepseek_max_output_tokens", 32000)
    monkeypatch.setattr(settings, "deepseek_request_timeout_seconds", 300)
    monkeypatch.setattr(settings, "deepseek_thinking_enabled", True)
    seen: list[int] = []

    class Completions:
        def create(self, **kwargs):
            seen.append(kwargs["max_tokens"])
            if kwargs["max_tokens"] > 8000:
                raise _rejection("max_tokens is too large: must be <= 8192")
            return SimpleNamespace(ok=True)

    client = SimpleNamespace(chat=SimpleNamespace(completions=Completions()))
    result = json_completion(client, [{"role": "user", "content": "hi"}], model="m")

    assert result.ok is True
    # 只降一次，且只在错误明确指向预算时降
    assert seen == [32000, 8000]


def test_completion_does_not_retry_unrelated_errors(monkeypatch):
    monkeypatch.setattr(settings, "deepseek_max_output_tokens", 32000)
    calls: list[dict] = []

    class Completions:
        def create(self, **kwargs):
            calls.append(kwargs)
            raise _rejection("invalid api key")

    client = SimpleNamespace(chat=SimpleNamespace(completions=Completions()))
    with pytest.raises(APIStatusError):
        json_completion(client, [], model="m")

    # 鉴权失败重试没有意义：重试只会把真实故障埋掉
    assert len(calls) == 1


def test_completion_turns_thinking_off_when_the_provider_rejects_it(monkeypatch):
    monkeypatch.setattr(settings, "deepseek_max_output_tokens", 8192)
    monkeypatch.setattr(settings, "deepseek_thinking_enabled", True)
    bodies: list[dict] = []

    class Completions:
        def create(self, **kwargs):
            bodies.append(kwargs["extra_body"])
            if kwargs["extra_body"]["thinking"]["type"] == "enabled":
                raise _rejection("thinking is not supported for this model")
            return SimpleNamespace(ok=True)

    client = SimpleNamespace(chat=SimpleNamespace(completions=Completions()))
    result = json_completion(client, [], model="m")

    assert result.ok is True
    assert bodies[-1] == {"thinking": {"type": "disabled"}}


# ── 深度思考开关：从请求一路到 provider ──────────────────


def test_plan_build_request_keeps_the_thinking_switch_three_state():
    """None 表示"按服务端配置"：老客户端不发这个字段时行为完全不变。"""
    assert CoursewarePlanBuildRequest().deep_thinking is None
    assert CoursewarePlanBuildRequest(deep_thinking=True).deep_thinking is True
    assert CoursewarePlanBuildRequest(deep_thinking=False).deep_thinking is False


def test_thinking_switch_overrides_the_server_default(monkeypatch):
    monkeypatch.setattr(settings, "deepseek_thinking_enabled", False)
    bodies: list[str] = []

    class Completions:
        def create(self, **kwargs):
            bodies.append(kwargs["extra_body"]["thinking"]["type"])
            return SimpleNamespace(ok=True)

    client = SimpleNamespace(chat=SimpleNamespace(completions=Completions()))
    for switch in (True, False, None):
        _completion(client, [{"role": "user", "content": "hi"}], thinking=switch)

    # 显式开关压过服务端默认；None 回落默认
    assert bodies == ["enabled", "disabled", "disabled"]


def test_completer_binds_the_switch_for_every_call(monkeypatch):
    """整条流水线的每一次补全都走同一个绑定好的 completer。

    逐个调用点传参迟早会漏一处，而最容易漏的正好是修复与审校步骤——所以绑定只做一次。
    """
    recorded: list[bool | None] = []

    def fake_completion(_client, _messages, *, thinking=None):
        recorded.append(thinking)
        return SimpleNamespace(ok=True)

    monkeypatch.setattr(courseware_ai, "_completion", fake_completion)
    complete = _make_completer(object(), False)

    complete([])
    complete([])

    assert recorded == [False, False]


def test_pipeline_binds_the_switch_from_its_argument(monkeypatch):
    """生成入口收到的 deep_thinking 必须真的绑定进去，而不是被丢掉。"""
    recorded: list[bool | None] = []

    def factory(_client, thinking):
        recorded.append(thinking)

        def complete(_messages):
            raise RuntimeError("bound")

        return complete

    monkeypatch.setattr(courseware_ai, "_make_completer", factory)

    with pytest.raises(RuntimeError):
        generate_courseware_spec_agentic(
            BRIEF_CONTENT, [], client=object(), thinking=False
        )

    assert recorded == [False]


# ── 模板兜底路径：必须是"能照着上"的内容，而不是占位句 ────

TEMPLATE_BRIEF = {
    "teaching_goal": "理解 TCP 三次握手",
    "subject": "信息技术",
    "grade": "大学",
    "target_audience": "计算机专业大二学生",
    "duration_minutes": 45,
    "knowledge_points": [
        {
            "order": 1,
            "title": "连接建立过程",
            "difficulty": "basic",
            "key_points": ["客户端发送 SYN", "服务器返回 SYN-ACK", "客户端确认 ACK"],
            "examples": ["客户端与服务器建立连接"],
            "estimated_minutes": 20,
        },
        {
            "order": 2,
            "title": "报文确认机制",
            "difficulty": "intermediate",
            "key_points": ["序列号", "确认号"],
            "estimated_minutes": 15,
        },
    ],
    "logic_flow": ["问题导入", "过程讲解", "例题练习", "总结"],
    "teaching_focus": "三次报文交换的时序和作用",
    "teaching_difficulties": "区分 SYN 与 ACK 的含义",
    "output_types": ["pptx", "docx", "pdf", "html"],
    "forbidden_content": "不要涉及微积分",
    "extra_requirements": "每个环节都要有板书设计",
    "homework_type": "绘制三次握手时序图",
}

TEMPLATE_REFS = [
    EvidenceRef(
        evidence_id="evidence-template-1",
        source_type="material",
        source_name="三次握手课程资料.pdf",
        quote="三次握手通过 SYN、SYN-ACK 和 ACK 三个报文建立连接。",
    )
]


def _template_brief(**overrides) -> SimpleNamespace:
    return SimpleNamespace(content_json={**TEMPLATE_BRIEF, **overrides})


def test_template_fallback_meets_the_richness_floors():
    """模板是确定性兜底，但同样必须是"能照着上"的内容。

    它以前只把 brief 拼成短句：讲稿 20 余字、每个环节师生活动各 1 条、
    打印要点一句话——质量检查会如实报"内容单薄"。
    """
    plan = compile_plan_content(_template_brief(), TEMPLATE_REFS)
    report = inspect_courseware(plan)

    assert report["status"] == "passed", report
    metrics = report["metrics"]
    assert metrics["speaker_notes_chars_min"] >= MIN_SPEAKER_NOTES_CHARS
    assert metrics["printable_summary_chars"] >= MIN_PRINTABLE_SUMMARY_CHARS
    assert metrics["checklist_items"] >= MIN_CHECKLIST_ITEMS
    # 每个环节都要有可执行的师生活动，而不是"学生完成练习"这种一句话
    for section in plan.lesson_sections:
        assert len(section.teacher_actions) >= 2
        assert len(section.student_actions) >= 2
        assert "典型错误" in section.assessment


def test_template_fallback_honours_explicit_teacher_requirements():
    """教师提的约束与额外要求要真的进产物，而不是只躺在需求单里。"""
    plan = compile_plan_content(_template_brief(), TEMPLATE_REFS)
    preparation = " ".join(plan.output_specs.docx.teacher_preparation)

    # 约束以拆分后的词条呈现，而不是把整句否定话术原样抄进教案
    assert "微积分" in preparation
    assert "不要涉及微积分" not in preparation
    assert "每个环节都要有板书设计" in preparation
    # 作业形式照教师给的写，并补上完成标准与提交形式
    assert plan.output_specs.docx.homework.startswith("绘制三次握手时序图")
    assert "完成标准" in plan.output_specs.docx.homework


def test_template_classification_lists_its_group_labels():
    """题干必须列出全部分组标签：只说"归类"而答案里有两组，学生无从作答。"""
    plan = compile_plan_content(
        _template_brief(
            teaching_goal="认识浮力",
            logic_flow=["导入", "讲解", "练习", "总结"],
            interaction_ideas="",
            knowledge_points=[
                {"order": 1, "title": "浮力的方向", "difficulty": "basic", "estimated_minutes": 20},
                {"order": 2, "title": "阿基米德原理", "difficulty": "intermediate", "estimated_minutes": 25},
            ],
        ),
        TEMPLATE_REFS,
    )
    interaction = plan.interactions[0]

    assert interaction.interaction_type == "classification"
    labels = list(interaction.answer_groups)
    assert len(labels) >= 2  # basic / intermediate 两个难度
    for label in labels:
        assert label in interaction.prompt
    assert interaction.explanation.strip()
