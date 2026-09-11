"""Legacy text-only reference parser for PDF and DOCX files."""
import re
import logging
import os

import fitz  # PyMuPDF
from docx import Document

# 初始化日志记录器
logger = logging.getLogger(__name__)

async def parse_pdf(file_path: str) -> str:
    """Extract structured text from PDF; embedded images are intentionally ignored."""
    if not os.path.exists(file_path):
        logger.error(f"PDF 文件不存在: {file_path}")
        raise FileNotFoundError(f"文件不存在: {file_path}")

    doc = fitz.open(file_path)
    full_content = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        full_content.append(f"\n\n=========== 第 {page_num + 1} 页 ===========\n")

        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            # 👑 针对痛点1：根据字号大小，动态赋予 Markdown 标题层级，防止结构丢失
            if block["type"] == 0:
                block_text = ""
                max_font_size = 0
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        block_text += span.get("text", "")
                        if span.get("size", 0) > max_font_size:
                            max_font_size = span.get("size", 0)
                    block_text += "\n"

                block_text = block_text.strip()
                if not block_text:
                    continue

                if max_font_size >= 15:
                    full_content.append(f"\n# {block_text}\n\n")
                elif max_font_size >= 12:
                    full_content.append(f"\n## {block_text}\n\n")
                else:
                    full_content.append(f"{block_text}\n\n")

    doc.close()
    return "".join(full_content).strip()

async def parse_docx(file_path: str) -> str:
    """Extract text and tables from DOCX; embedded images are intentionally ignored."""
    if not os.path.exists(file_path):
        logger.error(f"Word 文件不存在: {file_path}")
        raise FileNotFoundError(f"文件不存在: {file_path}")

    doc = Document(file_path)
    content_blocks = []

    # 1. 提取文字排版 (保留列表和缩进)
    for para in doc.paragraphs:
        text = para.text.rstrip()
        if not text.strip():
            continue

        if not text.startswith(" ") and not text.startswith("\t"):
            has_indent = False
            try:
                fmt = para.paragraph_format
                if (fmt.first_line_indent and fmt.first_line_indent > 0) or \
                   (fmt.left_indent and fmt.left_indent > 0):
                    has_indent = True
            except Exception:
                pass
            if has_indent:
                text = "    " + text

        clean_text = text.strip()
        if len(content_blocks) == 0:
            content_blocks.append(f"{text}\n\n")
        elif re.match(r'^[一二三四五六七八九十]+、', clean_text):
            content_blocks.append(f"\n### {text}\n")
        else:
            content_blocks.append(f"{text}\n")

    # 👑 针对痛点1：提取表格，并强制转换为标准的 Markdown 表格语法
    if doc.tables:
        for table in doc.tables:
            content_blocks.append("\n**【文档表格数据】**\n")
            for i, row in enumerate(table.rows):
                row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                # 拼接成 | 单元格 | 单元格 | 的格式
                content_blocks.append("| " + " | ".join(row_data) + " |\n")
                if i == 0:
                    # 第一行结束后，加上 Markdown 的表头分割线
                    content_blocks.append("|" + "|".join(["---"] * len(row.cells)) + "|\n")
            content_blocks.append("\n")

    return "".join(content_blocks).strip()
