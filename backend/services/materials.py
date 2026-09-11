"""Material ingestion and deterministic local parsing for the M3 workflow."""

from __future__ import annotations

import hashlib
import io
import mimetypes
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from PIL import Image
from pptx import Presentation

from backend.config import settings
from backend.db.database import SessionLocal
from backend.models.material import EvidenceChunk, Material, MaterialAnalysis
from video_parser.parser import PARSER_VERSION as VIDEO_PARSER_VERSION
from video_parser.parser import parse_video as video_parse_video
from video_parser.schemas import VideoParseOptions, VideoParseResult


class MaterialValidationError(ValueError):
    """Raised when an uploaded file fails type or content validation."""

    def __init__(self, message: str, *, code: str, details: dict | None = None):
        super().__init__(message)
        self.code = code
        self.details = details or {}


@dataclass(frozen=True)
class ParsedChunk:
    text: str
    locator: dict
    metadata: dict


@dataclass(frozen=True)
class ParsedMaterial:
    text_content: str
    chunks: list[ParsedChunk]
    result_json: dict
    page_count: int | None = None
    slide_count: int | None = None
    duration_seconds: int | None = None


FILE_TYPE_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "word",
    ".pptx": "ppt",
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
    ".gif": "image",
    ".bmp": "image",
    ".webp": "image",
    ".mp4": "video",
    ".avi": "video",
    ".mov": "video",
    ".mkv": "video",
}

PARSER_VERSION = "builtin-1"
VIDEO_MATERIAL_PARSER_NAME = "video-parser-model"
VIDEO_MATERIAL_PARSER_VERSION = VIDEO_PARSER_VERSION
MAX_CHUNK_CHARS = 1200


def checksum_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def detect_file_type(filename: str, content: bytes) -> tuple[str, str]:
    """Validate the extension against a small set of content signatures."""
    extension = Path(filename).suffix.lower()
    file_type = FILE_TYPE_EXTENSIONS.get(extension)
    if file_type is None:
        raise MaterialValidationError(
            f"不支持的文件格式: {extension or 'unknown'}",
            code="UNSUPPORTED_MATERIAL_FORMAT",
            details={"extension": extension},
        )

    if file_type == "pdf" and not content.startswith(b"%PDF-"):
        raise MaterialValidationError(
            "文件扩展名与真实 PDF 类型不一致",
            code="MATERIAL_TYPE_MISMATCH",
        )

    if file_type in {"word", "ppt"}:
        required_member = "word/document.xml" if file_type == "word" else "ppt/presentation.xml"
        if not _zip_contains(content, required_member):
            raise MaterialValidationError(
                "Office 文件内容无效或类型不匹配",
                code="MATERIAL_TYPE_MISMATCH",
            )

    if file_type == "image":
        try:
            with Image.open(io.BytesIO(content)) as image:
                image.verify()
        except Exception as exc:
            raise MaterialValidationError(
                "图片内容无法验证",
                code="MATERIAL_TYPE_MISMATCH",
            ) from exc
    if file_type == "video" and not _looks_like_video(content):
        raise MaterialValidationError(
            "视频文件内容无效或类型不匹配",
            code="MATERIAL_TYPE_MISMATCH",
        )

    mime_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    return file_type, mime_type


def detect_file_type_from_path(filename: str, path: Path) -> tuple[str, str]:
    """Validate a staged upload without loading the complete file into memory."""
    extension = Path(filename).suffix.lower()
    file_type = FILE_TYPE_EXTENSIONS.get(extension)
    if file_type is None:
        raise MaterialValidationError(
            f"不支持的文件格式: {extension or 'unknown'}",
            code="UNSUPPORTED_MATERIAL_FORMAT",
            details={"extension": extension},
        )

    if file_type == "pdf":
        with path.open("rb") as source:
            if source.read(5) != b"%PDF-":
                raise MaterialValidationError(
                    "文件扩展名与真实 PDF 类型不一致",
                    code="MATERIAL_TYPE_MISMATCH",
                )
    elif file_type in {"word", "ppt"}:
        required_member = "word/document.xml" if file_type == "word" else "ppt/presentation.xml"
        try:
            with zipfile.ZipFile(path) as archive:
                valid = required_member in archive.namelist()
        except zipfile.BadZipFile:
            valid = False
        if not valid:
            raise MaterialValidationError(
                "Office 文件内容无效或类型不匹配",
                code="MATERIAL_TYPE_MISMATCH",
            )
    elif file_type == "image":
        try:
            with Image.open(path) as image:
                image.verify()
        except Exception as exc:
            raise MaterialValidationError(
                "图片内容无法验证",
                code="MATERIAL_TYPE_MISMATCH",
            ) from exc
    elif file_type == "video":
        with path.open("rb") as source:
            header = source.read(12)
        if not _looks_like_video(header):
            raise MaterialValidationError(
                "视频文件内容无效或类型不匹配",
                code="MATERIAL_TYPE_MISMATCH",
            )

    mime_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    return file_type, mime_type


def parse_material(file_type: str, path: Path) -> ParsedMaterial:
    """Parse supported material formats into searchable, locatable evidence."""
    if file_type == "pdf":
        return _parse_pdf(path)
    if file_type == "word":
        return _parse_docx(path)
    if file_type == "ppt":
        return _parse_pptx(path)
    if file_type == "image":
        return _parse_image(path)
    if file_type == "video":
        return _parse_video(path)
    raise MaterialValidationError("暂不支持该资料类型", code="UNSUPPORTED_MATERIAL_TYPE")


def parser_identity(file_type: str) -> tuple[str, str]:
    if file_type == "video":
        return VIDEO_MATERIAL_PARSER_NAME, VIDEO_MATERIAL_PARSER_VERSION
    return f"builtin_{file_type}", PARSER_VERSION


def apply_parsed_material(
    db,
    material: Material,
    analysis: MaterialAnalysis,
    parsed: ParsedMaterial,
) -> None:
    completed_at = datetime.now(timezone.utc)
    analysis.status = "completed"
    analysis.text_content = parsed.text_content
    analysis.result_json = parsed.result_json
    analysis.page_count = parsed.page_count
    analysis.slide_count = parsed.slide_count
    analysis.duration_seconds = parsed.duration_seconds
    analysis.error_code = None
    analysis.error_message = None
    analysis.completed_at = completed_at
    analysis.updated_at = completed_at
    material.status = "ready"
    material.error_code = None
    material.error_message = None
    material.updated_at = completed_at
    for index, chunk in enumerate(parsed.chunks):
        db.add(
            EvidenceChunk(
                evidence_id=f"evidence_{uuid.uuid4().hex[:24]}",
                material_id=material.material_id,
                analysis_id=analysis.analysis_id,
                source_type=f"uploaded_{material.file_type}",
                chunk_index=index,
                locator_json=chunk.locator,
                text=chunk.text,
                metadata_json=chunk.metadata,
                usage_tags=[],
                content_hash=checksum_sha256(chunk.text.encode("utf-8")),
                is_valid=True,
                created_at=completed_at,
            )
        )


def fail_material_analysis(
    db,
    material: Material,
    analysis: MaterialAnalysis,
    error: Exception | str,
    *,
    code: str = "MATERIAL_PARSE_FAILED",
    status: str = "failed",
) -> None:
    now = datetime.now(timezone.utc)
    message = str(error)
    material.status = status
    material.error_code = code
    material.error_message = message
    material.updated_at = now
    analysis.status = "pending" if status == "queued" else status
    analysis.error_code = code
    analysis.error_message = message
    if status == "failed":
        analysis.completed_at = now
    analysis.updated_at = now


def run_material_analysis(analysis_id: str, *, raise_errors: bool = False) -> str | None:
    db = SessionLocal()
    try:
        analysis = (
            db.query(MaterialAnalysis)
            .filter(MaterialAnalysis.analysis_id == analysis_id)
            .with_for_update()
            .first()
        )
        if analysis is None or analysis.status == "completed":
            return analysis_id if analysis else None
        material = (
            db.query(Material)
            .filter(Material.material_id == analysis.material_id)
            .with_for_update()
            .first()
        )
        if material is None or material.deleted_at is not None:
            return None
        now = datetime.now(timezone.utc)
        material.status = "processing"
        material.error_code = None
        material.error_message = None
        material.updated_at = now
        analysis.status = "processing"
        analysis.started_at = analysis.started_at or now
        analysis.completed_at = None
        analysis.error_code = None
        analysis.error_message = None
        analysis.updated_at = now
        db.commit()

        parsed = parse_material(material.file_type, Path(material.stored_path))

        analysis = db.query(MaterialAnalysis).filter(MaterialAnalysis.analysis_id == analysis_id).one()
        material = db.query(Material).filter(Material.material_id == analysis.material_id).one()
        db.query(EvidenceChunk).filter(EvidenceChunk.analysis_id == analysis.analysis_id).delete()
        apply_parsed_material(db, material, analysis, parsed)
        db.commit()
        return analysis_id
    except Exception as exc:
        db.rollback()
        analysis = db.query(MaterialAnalysis).filter(MaterialAnalysis.analysis_id == analysis_id).first()
        material = (
            db.query(Material).filter(Material.material_id == analysis.material_id).first()
            if analysis is not None
            else None
        )
        if analysis is not None and material is not None:
            fail_material_analysis(db, material, analysis, exc)
            db.commit()
        if raise_errors:
            raise
        return None
    finally:
        db.close()


def prepare_material_retry(analysis_id: str, error: Exception | str) -> bool:
    db = SessionLocal()
    try:
        analysis = db.query(MaterialAnalysis).filter(MaterialAnalysis.analysis_id == analysis_id).first()
        material = (
            db.query(Material).filter(Material.material_id == analysis.material_id).first()
            if analysis is not None
            else None
        )
        if analysis is None or material is None or analysis.status == "completed":
            return False
        fail_material_analysis(
            db,
            material,
            analysis,
            error,
            code="MATERIAL_PARSE_RETRYING",
            status="queued",
        )
        db.commit()
        return True
    finally:
        db.close()


def _zip_contains(content: bytes, member: str) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            return member in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _looks_like_video(content: bytes) -> bool:
    if len(content) < 12:
        return False
    if content.startswith(b"\x1a\x45\xdf\xa3"):
        return True
    if content.startswith(b"RIFF") and content[8:12] == b"AVI ":
        return True
    return content[4:8] == b"ftyp"


def _parse_pdf(path: Path) -> ParsedMaterial:
    chunks: list[ParsedChunk] = []
    pages: list[str] = []
    with fitz.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if not text:
                continue
            pages.append(text)
            chunks.extend(_chunks(text, {"page": page_number}, {"format": "pdf"}))
        page_count = len(document)
    return ParsedMaterial(
        text_content="\n\n".join(pages),
        chunks=chunks,
        result_json={"format": "pdf", "page_count": page_count, "chunk_count": len(chunks)},
        page_count=page_count,
    )


def _parse_docx(path: Path) -> ParsedMaterial:
    document = Document(path)
    chunks: list[ParsedChunk] = []
    blocks: list[str] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
            chunks.extend(_chunks(text, {"paragraph": index}, {"format": "docx"}))

    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells) for row in table.rows]
        table_text = "\n".join(row for row in rows if row.strip())
        if table_text:
            blocks.append(table_text)
            chunks.extend(
                _chunks(
                    table_text,
                    {"table": table_index},
                    {"format": "docx", "content_type": "table"},
                )
            )

    return ParsedMaterial(
        text_content="\n\n".join(blocks),
        chunks=chunks,
        result_json={"format": "docx", "paragraph_count": len(document.paragraphs), "chunk_count": len(chunks)},
    )


def _parse_pptx(path: Path) -> ParsedMaterial:
    presentation = Presentation(path)
    chunks: list[ParsedChunk] = []
    slides: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_blocks = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text")]
        text = "\n".join(block for block in text_blocks if block)
        if not text:
            continue
        slides.append(text)
        chunks.extend(_chunks(text, {"slide": slide_number}, {"format": "pptx"}))
    return ParsedMaterial(
        text_content="\n\n".join(slides),
        chunks=chunks,
        result_json={"format": "pptx", "slide_count": len(presentation.slides), "chunk_count": len(chunks)},
        slide_count=len(presentation.slides),
    )


def _parse_image(path: Path) -> ParsedMaterial:
    with Image.open(path) as image:
        metadata = {
            "format": image.format,
            "width": image.width,
            "height": image.height,
            "mode": image.mode,
            "requires_vision": True,
            "vision_status": "not_configured",
            "warning": "图片已保存，但视觉识别模型尚未配置，未生成证据文本。",
        }
    return ParsedMaterial(text_content="", chunks=[], result_json=metadata)


def _parse_video(path: Path) -> ParsedMaterial:
    options = VideoParseOptions(
        video_type=_video_type(),
        transcribe=settings.video_parser_transcribe,
        ocr=settings.video_parser_ocr,
        vision=settings.video_parser_vision,
        video_understanding=settings.video_parser_understanding,
        max_sample_keyframes=settings.video_parser_max_keyframes,
        max_shots=settings.video_parser_max_shots,
        output_width=settings.video_parser_output_width,
    )
    result = video_parse_video(path, output_root=settings.video_parser_output_dir, options=options)
    return _video_result_to_material(result)


def _video_type() -> str:
    if settings.video_parser_type in {"auto", "presentation", "whiteboard", "operation"}:
        return settings.video_parser_type
    return "auto"


def _video_result_to_material(result: VideoParseResult) -> ParsedMaterial:
    chunks = [_video_evidence_to_chunk(item) for item in result.evidence if item.content.strip()]
    segment_lines = [
        f"{segment.time_range.start}-{segment.time_range.end} {segment.summary}"
        for segment in result.segments
        if segment.summary.strip()
    ]
    text_parts = [
        result.transcript.text.strip(),
        *segment_lines,
        *(item.text for item in chunks if item.text.strip()),
    ]
    payload = result.model_dump(mode="json")
    return ParsedMaterial(
        text_content="\n".join(dict.fromkeys(item for item in text_parts if item)),
        chunks=chunks,
        result_json={
            "format": "video",
            "parser": {"name": VIDEO_MATERIAL_PARSER_NAME, "version": VIDEO_MATERIAL_PARSER_VERSION},
            "video_id": result.video_id,
            "duration_seconds": result.metadata.duration_seconds,
            "segment_count": len(result.segments),
            "keyframe_count": len(result.keyframes),
            "evidence_count": len(result.evidence),
            "warnings": result.warnings,
            "artifacts": result.artifacts,
            "result": payload,
        },
        duration_seconds=round(result.metadata.duration_seconds),
    )


def _video_evidence_to_chunk(item) -> ParsedChunk:
    time_range = item.time_range.model_dump(mode="json") if item.time_range else None
    locator: dict[str, Any] = {
        "source_id": item.source_id,
        "evidence_type": item.evidence_type,
    }
    if time_range:
        locator["timestamp"] = time_range.get("start")
        locator["time_range"] = time_range
    metadata = {
        "format": "video",
        "evidence_type": item.evidence_type,
        "source_path": item.source_path,
        **(item.metadata or {}),
    }
    return ParsedChunk(text=item.content, locator=locator, metadata=metadata)


def _chunks(text: str, locator: dict, metadata: dict) -> list[ParsedChunk]:
    normalized = " ".join(text.split())
    if not normalized:
        return []
    parts = [normalized[index : index + MAX_CHUNK_CHARS] for index in range(0, len(normalized), MAX_CHUNK_CHARS)]
    return [
        ParsedChunk(
            text=part,
            locator={**locator, "part": index + 1} if len(parts) > 1 else locator,
            metadata=metadata,
        )
        for index, part in enumerate(parts)
    ]
